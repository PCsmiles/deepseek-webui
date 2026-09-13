#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
DeepSeek 本地聊天 Web UI —— 后端服务
=====================================

作用:
  1. 把网页发来的对话请求转发给 DeepSeek 的 OpenAI 兼容接口 (/chat/completions)
  2. 把 DeepSeek 的流式响应(SSE)原样转成"我们自己的 SSE"推给浏览器,实现逐字输出
  3. 保管 API Key(永远不下发给浏览器)
  4. 存图片(data/images)、存对话备份(data/conversations)
  5. 代理查余额(这样 Key 也不用给浏览器)

启动:  python server.py          (会自动打开浏览器)
可选:  python server.py --port 8888 --no-browser
"""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import glob
import json
import os
import re
import secrets
import shutil
import signal
import socket
import subprocess
import sys
import threading
import time
import webbrowser
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
import uvicorn
from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

# 让中文在 Windows 控制台里不乱码、不因为编码问题崩掉
if hasattr(sys.stdout, "reconfigure"):
    with contextlib.suppress(Exception):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    with contextlib.suppress(Exception):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")


# --------------------------------------------------------------------------
# 路径
#   config.json 从三个地方找（先找到的算）：
#     1. 环境变量 JY_CONFIG 指的路径
#     2. 程序目录的上一级（打包版：%LOCALAPPDATA%\Jingyu\config.json
#        —— 放在程序目录外面，重装/升级都不会把配置和聊天记录冲掉）
#     3. 程序目录里（开发时就是仓库根目录）
#   数据目录（会话/图片/日志）跟着 config.json 的 data_dir 走，默认在程序目录下的 data/。
# --------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
VENDOR_DIR = STATIC_DIR / "vendor"


def _read_json(p: Path) -> Dict[str, Any]:
    # utf-8-sig：别人用记事本另存为时可能带上 BOM，带上也能正常读
    try:
        d = json.loads(Path(p).read_text(encoding="utf-8-sig"))
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def _find_config_file() -> Path:
    for cand in (os.environ.get("JY_CONFIG"),
                 str(ROOT.parent / "config.json"),
                 str(ROOT / "config.json")):
        if cand and Path(cand).is_file():
            return Path(cand)
    return ROOT / "config.json"          # 还没有也算它（保存时就写到这里）


CONFIG_FILE = _find_config_file()
_RAW_CONF: Dict[str, Any] = _read_json(CONFIG_FILE)

DATA_DIR = Path(
    os.environ.get("JY_DATA_DIR")
    or str(_RAW_CONF.get("data_dir") or "").strip()
    # 没写 data_dir 就按"配置文件在哪"来定：
    #   config.json 在程序目录外面（打包版：%LOCALAPPDATA%\Jingyu\config.json）
    #   → 数据就放它旁边，卸载程序不会连带删掉聊天记录
    #   开发时（config 在仓库里）→ 还是放程序目录下的 data\
    or (str(CONFIG_FILE.parent / "data") if CONFIG_FILE.parent != ROOT else str(ROOT / "data"))
).expanduser()
CONV_DIR = DATA_DIR / "conversations"
IMG_DIR = DATA_DIR / "images"
for _d in (DATA_DIR, CONV_DIR, IMG_DIR, VENDOR_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# 静默模式（由 launch_silent.py 启动时设置）：没有控制台，就把输出写进 data/server.log
if os.environ.get("DSUI_SILENT") == "1":
    with contextlib.suppress(Exception):
        _lf = DATA_DIR / "server.log"
        if _lf.exists() and _lf.stat().st_size > 2 * 1024 * 1024:   # 超过 2MB 就重开
            _lf.unlink()
        _log = open(_lf, "a", encoding="utf-8", buffering=1)
        _log.write(f"\n===== 启动 {time.strftime('%Y-%m-%d %H:%M:%S')} =====\n")
        sys.stdout = _log
        sys.stderr = _log

VERSION = "1.2.0"
DEFAULT_BASE_URL = "https://api.deepseek.com"
DEFAULT_BRAND = "deepseek"                  # 界面上的名字（打包给别人的那份用 config.json 改成"鲸语"）
DEFAULT_PRETTY_HOST = "deepseek.localhost"  # 好看网址（*.localhost 才是浏览器眼里的"安全上下文"）
DEFAULT_SYSTEM_PROMPT = "你是一个乐于助人、回答准确的中文 AI 助手。回答用简体中文，代码要能直接运行。"

# 内置模型清单（DeepSeek 官方）。vision=True 表示能读图(2026-09-11 实测:flash 能,pro 不能)。
# config.json 里给了 "models" 就用它，方便换成中转站的模型名。
DEFAULT_MODELS = [
    {
        "id": "deepseek-flash",
        "label": "deepseek-flash",
        "desc": "V4.1 Flash · 快 · 便宜 · 能看图",
        "vision": True,
    },
    {
        "id": "deepseek-v4-pro",
        "label": "deepseek-v4-pro",
        "desc": "更强 · 慢一些 · 不支持图片",
        "vision": False,
    },
]

# 只允许这种格式的 id,防止有人构造 ../ 之类的路径读到别的文件
SAFE_ID = re.compile(r"^[A-Za-z0-9_-]{1,64}$")
# 图片 id 形如 1a08ed50724-bf80b51d.png：只允许一个点，且必须是图片后缀
SAFE_IMG = re.compile(r"^[A-Za-z0-9_-]{1,72}\.(?:png|jpe?g|webp|gif)$", re.I)


# --------------------------------------------------------------------------
# 配置：把 config.json + 环境变量 + Claude Code 的 settings.json 合成一份
#
# 新格式（打包给别人的那份 / 网页设置向导写出来的就是这种）：
#   {
#     "brand": "鲸语",
#     "pretty_host": "jingyu.localhost",
#     "chat":  { "protocol": "openai",            // openai | anthropic
#                "base_url": "https://api.deepseek.com",
#                "api_key": "sk-xxx", "model": "deepseek-flash" },
#     "agent": { "base_url": "https://api.deepseek.com/anthropic",
#                "api_key": "sk-xxx", "model": "deepseek-flash[1M]",
#                "workspace": "C:\\Users\\me\\Documents\\鲸语工作区" },
#     "asr_key": "",          // 可选：语音转文字（硅基流动）
#     "memory_dir": ""        // 可选：长期记忆目录；留空 = 用本机终端那套
#   }
# 老的平铺格式（"api_key" / "base_url" / "model" 直接写在最外层）也照样认。
# --------------------------------------------------------------------------
def _claude_code_env() -> Dict[str, str]:
    """读 ~/.claude/settings.json 里 env 那一段（本机一直靠它拿 Key）"""
    cfg = _read_json(Path.home() / ".claude" / "settings.json")
    env = cfg.get("env") or {}
    return {k: v.strip() for k, v in env.items() if isinstance(v, str) and v.strip()}


def load_config() -> Dict[str, Any]:
    raw = dict(_RAW_CONF)
    cce = _claude_code_env()

    chat = dict(raw.get("chat") or {})
    agent = dict(raw.get("agent") or {})
    for k in ("base_url", "api_key", "model"):        # 老的平铺写法 → chat
        if not chat.get(k) and isinstance(raw.get(k), str):
            chat[k] = raw[k]

    chat["protocol"] = "anthropic" if str(chat.get("protocol") or "").lower() == "anthropic" else "openai"
    chat["base_url"] = str(chat.get("base_url") or DEFAULT_BASE_URL).rstrip("/")
    chat["model"] = str(chat.get("model") or "deepseek-flash")

    # --- Key 三级查找：环境变量 > config.json > Claude Code 的 settings.json ---
    key, source = str(chat.get("api_key") or "").strip(), ""
    if key:
        source = "config.json（网页设置里填的）"
    env_key = (os.environ.get("DEEPSEEK_API_KEY") or "").strip()
    if env_key:
        key, source = env_key, "环境变量 DEEPSEEK_API_KEY"
    elif not key:
        ck = cce.get("ANTHROPIC_AUTH_TOKEN") or cce.get("ANTHROPIC_API_KEY") or ""
        if ck:
            key, source = ck, "自动复用 Claude Code 的 Key (~/.claude/settings.json)"
    chat["api_key"] = key

    # --- 干活模式（Claude Code）自己那一套地址/Key ---
    # 留空就什么都不设：让 claude 去读它自己的 ~/.claude/settings.json（本机就是这种）
    agent = {
        "base_url": str(agent.get("base_url") or "").strip().rstrip("/"),
        "api_key": str(agent.get("api_key") or "").strip(),
        "model": str(agent.get("model") or "").strip(),
        "workspace": str(agent.get("workspace") or raw.get("workspace") or "").strip(),
    }

    conf: Dict[str, Any] = {
        "brand": str(raw.get("brand") or DEFAULT_BRAND).strip() or DEFAULT_BRAND,
        "pretty_host": str(raw.get("pretty_host") or DEFAULT_PRETTY_HOST).strip() or DEFAULT_PRETTY_HOST,
        "host": str(raw.get("host") or "127.0.0.1"),
        "port": int(raw.get("port") or 80),
        "temperature": float(raw.get("temperature") or 0.7),
        "max_tokens": int(raw.get("max_tokens") or 4096),
        "system_prompt": str(raw.get("system_prompt") or DEFAULT_SYSTEM_PROMPT),
        "asr_key": str(raw.get("asr_key") or "").strip(),
        "memory_dir": str(raw.get("memory_dir") or "").strip(),
        "chat": chat,
        "agent": agent,
        # 老代码直接读这几个顶层字段，留着省得各处都改
        "api_key": key, "key_source": source,
        "base_url": chat["base_url"], "model": chat["model"],
    }
    if isinstance(raw.get("models"), list) and raw["models"]:
        conf["models"] = raw["models"]
    return conf


CONFIG = load_config()
AVAILABLE_MODELS: List[Dict[str, Any]] = CONFIG.get("models") or DEFAULT_MODELS
_have_key = bool(CONFIG["chat"]["api_key"])


# --------------------------------------------------------------------------
# 把 DeepSeek 的报错翻译成人话
# --------------------------------------------------------------------------
def humanize_error(status: int, raw: str) -> tuple[str, str, str]:
    """返回 (错误代码, 给用户看的话, 怎么办的建议)"""
    txt = (raw or "").strip()
    low = txt.lower()

    if status == 401:
        return ("auth", "API Key 无效或已过期,DeepSeek 拒绝了这次请求。",
                "检查 config.json 里的 api_key,或环境变量 DEEPSEEK_API_KEY;"
                "也可以确认一下 ~/.claude/settings.json 里的 Key 是否还有效。")
    if status == 402 or "insufficient balance" in low or "余额不足" in txt:
        return ("balance", "账户余额不足,DeepSeek 拒绝了这个请求。",
                "去 platform.deepseek.com 充值。侧边栏右下角可以看到当前余额。")
    if status == 429:
        return ("rate", "请求太频繁,被限流了。",
                "等 10 秒再发;如果经常这样,说明当前账号并发受限。")
    if status == 400:
        if "context length" in low or "too long" in low or "token" in low and "exceed" in low:
            return ("context", "这段对话太长了,超出了模型一次能处理的长度。",
                    "点左上角「新建对话」另起一个,或删掉一些早期消息。")
        if "max_tokens" in low:
            return ("param", "max_tokens 这个数值不合法。",
                    "打开设置,把 max_tokens 调到 1 ~ 8192 之间。")
        if "model" in low and ("not exist" in low or "invalid" in low or "supported" in low):
            return ("model", "模型名不对,DeepSeek 不认识这个名字。",
                    "当前可用模型只有 deepseek-flash 和 deepseek-v4-pro。"
                    "注意不要加 [1M] 后缀(那个只在 Claude Code 用的 Anthropic 端点有效)。")
        return ("bad_request", "请求被 DeepSeek 拒绝了(参数或内容有问题)。",
                f"原始报错: {txt[:300]}")
    if status == 404:
        return ("notfound", "接口地址不对(404)。",
                f"检查 config.json 的 base_url,当前是 {CONFIG['base_url']}。")
    if status in (500, 502, 503, 504):
        return ("server", "DeepSeek 服务端出错了,不是你这边的问题。",
                "等 1 分钟重试;一直这样的话去 status.deepseek.com 看看是不是在维护。")
    return ("http", f"DeepSeek 返回了 HTTP {status}。", f"原始报错: {txt[:300]}")


def network_error(exc: Exception) -> tuple[str, str, str]:
    name = type(exc).__name__
    if isinstance(exc, httpx.ConnectTimeout):
        return ("timeout_connect", "连不上 api.deepseek.com(连接超时)。",
                "如果开着代理/VPN,试试关掉——DeepSeek 是国内服务,直连最快。")
    if isinstance(exc, httpx.ReadTimeout):
        return ("timeout_read", "等回答等超时了(模型太久没吐字)。",
                "直接点「重新生成」;如果反复超时,把 max_tokens 调小一点。")
    if isinstance(exc, httpx.ConnectError):
        return ("network", "网络连不通(可能是断网或 DNS 问题)。",
                "浏览器里打开 https://api.deepseek.com 试试能不能通。")
    if isinstance(exc, httpx.RemoteProtocolError):
        return ("protocol", "连接被中途掐断了。", "重试一次通常就好。")
    return ("unknown", f"调用时出了点意外({name})。", f"{exc}")


# --------------------------------------------------------------------------
# FastAPI 应用
# --------------------------------------------------------------------------
@asynccontextmanager
async def lifespan(app: FastAPI):
    app.state.client = httpx.AsyncClient(
        timeout=httpx.Timeout(connect=15.0, read=600.0, write=60.0, pool=15.0),
        limits=httpx.Limits(max_connections=32, max_keepalive_connections=16),
    )
    app.state.balance_cache = {"ts": 0.0, "data": None}
    print("=" * 62)
    print(f"  {CONFIG['brand']} 本地工作台  v{VERSION}")
    print(f"  项目目录 : {ROOT}")
    print(f"  配置文件 : {CONFIG_FILE}")
    print(f"  数据目录 : {DATA_DIR}")
    print(f"  聊天接口 : {CONFIG['chat']['base_url']}  [{CONFIG['chat']['protocol']}]  模型 {CONFIG['chat']['model']}")
    print(f"  聊天 Key : {'已就绪 <- ' + CONFIG['key_source'] if _have_key else '*** 没有找到 API Key ***'}")
    _ag = CONFIG["agent"]
    print(f"  干活接口 : {_ag['base_url'] or '(用 Claude Code 自己的配置)'}  "
          f"{'Key 已就绪' if _ag['api_key'] else ''}  模型 {_ag['model'] or '(跟随 Claude Code 配置)'}")
    print(f"  claude   : {CLAUDE_BIN or '*** 没找到 claude 命令 ***'}")
    print("=" * 62)
    if not _have_key:
        print("[!] 还没有配 API Key —— 打开网页会直接弹出设置向导，或者：")
        print("    1) 在 config.json 里填 \"chat\": {\"api_key\": \"sk-xxxxxx\"}")
        print("    2) 设置环境变量 DEEPSEEK_API_KEY")
        print("    3) 确认 ~/.claude/settings.json 里有 ANTHROPIC_AUTH_TOKEN")
    # 打包版第一次跑的时候，把手配的工作目录先建出来（不然要等第一次干活才建）
    _ws = CONFIG["agent"]["workspace"]
    if _ws:
        with contextlib.suppress(Exception):
            Path(_ws).mkdir(parents=True, exist_ok=True)
    yield
    await app.state.client.aclose()


class NoCacheStaticFiles(StaticFiles):
    """静态文件永远"回源校验"，不许浏览器自己拿缓存里的旧副本。

    默认的 StaticFiles 不发 Cache-Control，浏览器就按"启发式缓存"自己算有效期
    （≈ 文件年龄的 10%，可能是好几个小时）。于是会出现：index.html 已经换成新版，
    app.js 还在用缓存里的旧版 → 旧代码找不到新页面里的元素 →
    界面报 "Cannot read properties of null"（用户 2026-09-13 报过，F5 才能好）。
    加了 no-cache 后每次都带 ETag 回源问一句，没改就 304，本地几乎零成本。
    """

    async def get_response(self, path, scope):
        resp = await super().get_response(path, scope)
        resp.headers["Cache-Control"] = "no-cache"
        return resp


app = FastAPI(title="DeepSeek 本地聊天 UI", lifespan=lifespan)
app.mount("/static", NoCacheStaticFiles(directory=str(STATIC_DIR)), name="static")

# --------------------------------------------------------------------------
# 防盗用令牌
#   现在网页能让 AI 在你电脑上跑命令了，必须防止"你随便打开一个恶意网页，
#   那个网页偷偷往 127.0.0.1:8765 发请求让你的电脑干活"。
#   做法：所有写操作都要带一个随机令牌（只有本地这个页面知道，别的网站读不到），
#        再顺手校验 Origin，双保险。
# --------------------------------------------------------------------------
# 令牌要**持久化**：以前每次重启都换新的，害得已经打开的页面全部失效
# （用户实测报过"访问令牌不对"）。现在存在 data/token.txt，重启后不变。
def _load_token() -> str:
    env = (os.environ.get("DSUI_TOKEN") or "").strip()
    if env:
        return env
    f = DATA_DIR / "token.txt"
    with contextlib.suppress(Exception):
        old = f.read_text(encoding="utf-8").strip()
        if len(old) >= 16:
            return old
    tok = secrets.token_urlsafe(24)
    with contextlib.suppress(Exception):
        f.write_text(tok, encoding="utf-8")
    return tok


TOKEN = _load_token()
# 允许的来源：本机各种写法 + 任何 *.localhost（deepseek.localhost 也在这里面）。
# ⚠️ 千万别漏掉 .localhost：换了好看的域名却没更新这里，会导致浏览器发的请求全被 403
# （而且 Python 直连测试不带 Origin 头，测不出来 —— 踩过一次）
_ORIGIN_OK = re.compile(
    r"^http://(127\.0\.0\.1|localhost|\[::1\]|[A-Za-z0-9._-]+\.localhost)(:\d+)?$"
)


@app.middleware("http")
async def _token_guard(request: Request, call_next):
    if request.method not in ("GET", "HEAD", "OPTIONS") and request.url.path.startswith("/api/"):
        tok = request.headers.get("x-dsui-token", "")
        ok = bool(tok) and secrets.compare_digest(tok, TOKEN)
        org = request.headers.get("origin", "")
        if org and not _ORIGIN_OK.match(org):
            ok = False
        if not ok:
            return JSONResponse(
                {"ok": False, "message": "访问令牌不对（防止别的网页偷用你的电脑）"},
                status_code=403)
    return await call_next(request)


@app.get("/")
async def index():
    f = STATIC_DIR / "index.html"
    if not f.exists():
        return JSONResponse({"error": "static/index.html 不见了,文件不完整"},
                            status_code=500)
    return _index_response()


def _index_response() -> Response:
    """发 index.html，并给里面的 app.js / style.css 自动加上版本号（取文件修改时间）。

    为什么必须这么做：浏览器缓存按"网址"存。网址不变，它就可能拿旧的 app.js；
    而 index.html 本身改了（比如加/删了某个元素）→ 出现"新页面 + 旧脚本"，
    脚本找不到元素就报 Cannot read properties of null，整页卡住，F5 才好。
    现在代码一改，?v= 跟着变 → 浏览器被逼着重新下载，新旧错配从根上不可能发生。
    """
    newest = 0.0
    for name in ("app.js", "style.css", "index.html"):
        with contextlib.suppress(Exception):
            newest = max(newest, (STATIC_DIR / name).stat().st_mtime)
    ver = str(int(newest)) or "0"
    html = (STATIC_DIR / "index.html").read_text(encoding="utf-8")
    html = re.sub(r'((?:src|href)="/static/(?:app\.js|style\.css))"',
                  rf'\1?v={ver}"', html)
    return Response(html, media_type="text/html; charset=utf-8",
                    headers={"Cache-Control": "no-cache, must-revalidate"})


@app.get("/favicon.ico")
async def favicon():
    f = STATIC_DIR / "favicon.svg"
    if f.exists():
        return FileResponse(str(f), media_type="image/svg+xml")
    return JSONResponse({}, status_code=404)


def _reload_config() -> None:
    """改完 config.json 后重新加载（CONFIG 原地更新，别处拿着的引用照样有效）"""
    global _RAW_CONF, AVAILABLE_MODELS, _have_key
    _RAW_CONF = _read_json(CONFIG_FILE)
    fresh = load_config()
    CONFIG.clear()
    CONFIG.update(fresh)
    AVAILABLE_MODELS = CONFIG.get("models") or DEFAULT_MODELS
    _have_key = bool(CONFIG["chat"]["api_key"])


def _save_raw_conf(updates: Dict[str, Any]) -> Dict[str, Any]:
    """把 updates 合并进 config.json（chat / agent 这两段是深层合并）"""
    cur = _read_json(CONFIG_FILE)
    for k, v in updates.items():
        if isinstance(v, dict) and isinstance(cur.get(k), dict):
            cur[k] = {**cur[k], **v}
        else:
            cur[k] = v
    for sec in ("chat", "agent"):        # 别在文件里留一堆空字符串
        if isinstance(cur.get(sec), dict):
            cur[sec] = {k: v for k, v in cur[sec].items() if v not in (None, "")}
            if not cur[sec]:
                cur.pop(sec)
    with contextlib.suppress(Exception):
        CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
        CONFIG_FILE.write_text(json.dumps(cur, ensure_ascii=False, indent=2), encoding="utf-8")
    return cur


@app.get("/api/config")
async def api_config():
    """告诉前端:品牌、有哪些模型、Key 有没有准备好、干活模式能不能用。"""
    ag = CONFIG["agent"]
    return {
        "version": VERSION,
        "brand": CONFIG["brand"],
        "token": TOKEN,          # 前端拿它来带请求头；跨域网站读不到这个响应
        "agent_ready": bool(CLAUDE_BIN),
        "agent_configured": bool(ag["base_url"] and ag["api_key"]) or bool(_claude_code_env().get("ANTHROPIC_BASE_URL")),
        "agent_model": ag["model"],
        "agent_base_url": ag["base_url"],
        "claude_bin": CLAUDE_BIN,
        "models": AVAILABLE_MODELS,
        "default_model": CONFIG["model"],
        "protocol": CONFIG["chat"]["protocol"],
        "temperature": CONFIG["temperature"],
        "max_tokens": CONFIG["max_tokens"],
        "system_prompt": CONFIG["system_prompt"],
        "has_key": _have_key,
        "key_source": CONFIG["key_source"],
        "base_url": CONFIG["base_url"],
        "workspace": ag["workspace"],
        "data_dir": str(DATA_DIR),
        "config_file": str(CONFIG_FILE),
        "memory_enabled": bool(CONFIG["memory_dir"]) and Path(CONFIG["memory_dir"]).is_dir(),
        "vendor_ready": (VENDOR_DIR / "marked.min.js").exists(),
        # 附件能力自检，好让界面直接告诉用户"哪些能用"
        "attach": {
            "max_chars": MAX_INJECT_CHARS,
            "max_images": MAX_ATTACH_IMAGES,
            "asr": bool(_asr_key()),          # 语音转文字可用吗
            "ffmpeg": bool(_ffmpeg()),        # 视频处理可用吗
        },
    }


@app.post("/api/settings")
async def api_settings(request: Request):
    """写设置到 config.json（Key 永远不下发给浏览器，浏览器只能知道"配没配好"）。

    两种写法都收：
      · 老的平铺：{api_key, base_url, clear_key}          ← 设置面板里的快捷修改
      · 新的分段：{chat:{...}, agent:{...}, asr_key, brand} ← 首启向导
    """
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)

    upd: Dict[str, Any] = {}
    changed: List[str] = []

    def put(sec: str, k: str, v: Any) -> None:
        upd.setdefault(sec, {})[k] = v

    # —— 老写法 ——
    if isinstance(body.get("base_url"), str) and body["base_url"].strip():
        put("chat", "base_url", body["base_url"].strip())
        changed.append("接口地址已更新")
    if body.get("clear_key"):
        put("chat", "api_key", "")
        changed.append("已清除聊天 Key")
    elif isinstance(body.get("api_key"), str) and body["api_key"].strip():
        put("chat", "api_key", body["api_key"].strip())
        changed.append("聊天 Key 已更新")
        if (os.environ.get("DEEPSEEK_API_KEY") or "").strip():
            changed.append("⚠ 但环境变量 DEEPSEEK_API_KEY 优先级更高，会盖过这里")
    if body.get("clear_agent_key"):
        put("agent", "api_key", "")
        changed.append("已清除干活 Key")

    # —— 新写法（分段）——
    for sec, label in (("chat", "聊天"), ("agent", "干活")):
        patch = body.get(sec)
        if isinstance(patch, dict):
            clean = {k: (v.strip() if isinstance(v, str) else v)
                     for k, v in patch.items() if isinstance(v, (str, int, float))}
            clean = {k: v for k, v in clean.items() if v != ""}    # 空 = 不改
            if clean:
                upd.setdefault(sec, {}).update(clean)
                changed.append(f"{label}设置已更新")
    for k in ("asr_key", "brand", "memory_dir", "pretty_host"):
        if isinstance(body.get(k), str) and body[k].strip():
            upd[k] = body[k].strip()
            changed.append(f"{k} 已更新")
    if isinstance(body.get("models"), list) and body["models"]:
        upd["models"] = body["models"]
        changed.append("模型清单已更新")

    _save_raw_conf(upd)
    _reload_config()
    print(f"[settings] {'；'.join(changed) or '无变化'} -> has_key={_have_key}")
    return {"ok": True, "message": "；".join(changed) or "无变化",
            "has_key": _have_key, "key_source": CONFIG["key_source"],
            "base_url": CONFIG["base_url"], "protocol": CONFIG["chat"]["protocol"],
            "model": CONFIG["model"], "workspace": CONFIG["agent"]["workspace"]}


def _anthropic_url(base: str) -> str:
    """把用户填的地址拼成 Anthropic 的 /v1/messages（填了 /v1 结尾也不会拼成 /v1/v1）"""
    b = (base or "").rstrip("/")
    if b.endswith("/v1"):
        b = b[:-3]
    return b + "/v1/messages"


def _guess_anthropic_base(chat_base: str) -> str:
    """只知道聊天地址（OpenAI 那边）时，猜一下干活模式（Anthropic）该用哪个地址。
       DeepSeek 官方是同一个域名 + /anthropic；已经是 /anthropic 结尾的就别再加一遍。"""
    b = (chat_base or "").rstrip("/")
    if not b:
        return ""
    if b.endswith("/anthropic"):
        return b
    if "api.deepseek.com" in b:
        return b + "/anthropic"
    return ""


def _to_anthropic_messages(msgs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """把我们内部的（OpenAI 形状的）消息转成 Anthropic 的 content blocks"""
    out: List[Dict[str, Any]] = []
    for m in msgs:
        c = m.get("content")
        if isinstance(c, str):
            out.append({"role": m["role"], "content": [{"type": "text", "text": c}]})
            continue
        blocks: List[Dict[str, Any]] = []
        for p in (c or []):
            if not isinstance(p, dict):
                continue
            if p.get("type") == "text" and p.get("text"):
                blocks.append({"type": "text", "text": p["text"]})
            elif p.get("type") == "image_url":
                url = ((p.get("image_url") or {}).get("url") or "")
                mm = re.match(r"^data:(image/[A-Za-z0-9.+-]+);base64,(.+)$", url, re.S)
                if mm:
                    blocks.append({"type": "image", "source": {
                        "type": "base64", "media_type": mm.group(1), "data": mm.group(2)}})
        if blocks:
            out.append({"role": m["role"], "content": blocks})
    # Anthropic 要求 user/assistant 交替出现，连续的同一角色要合并
    merged: List[Dict[str, Any]] = []
    for m in out:
        if merged and merged[-1]["role"] == m["role"]:
            merged[-1]["content"] = merged[-1]["content"] + m["content"]
        else:
            merged.append(m)
    return merged


async def _pump_anthropic(j: Dict[str, Any], queue: "asyncio.Queue", acc: Dict[str, Any]) -> None:
    """把 Anthropic 的事件翻译成前端要的那几种（和 OpenAI 分支输出保持一模一样）"""
    t = j.get("type")
    if t == "content_block_delta":
        d = j.get("delta") or {}
        if d.get("type") == "thinking_delta" and d.get("thinking"):
            await queue.put({"type": "reasoning", "text": d["thinking"]})
        elif d.get("text"):
            await queue.put({"type": "content", "text": d["text"]})
    elif t == "message_start":
        u = ((j.get("message") or {}).get("usage")) or {}
        if u.get("input_tokens"):
            acc["prompt_tokens"] = u["input_tokens"]
    elif t == "message_delta":
        u = j.get("usage") or {}
        if u.get("output_tokens"):
            acc["completion_tokens"] = u["output_tokens"]
        sr = (j.get("delta") or {}).get("stop_reason")
        if sr:
            await queue.put({"type": "finish", "reason": "length" if sr == "max_tokens" else sr})
    elif t == "error":
        e = j.get("error") or {}
        await queue.put({"type": "error", "code": "upstream",
                         "message": f"接口报错：{e.get('message') or '未知'}",
                         "hint": str(e.get("type") or "")[:200]})
    # Anthropic 的输入/输出 token 是分两次来的，攒成一条 usage 发（前端直接显示总数）
    if t in ("message_delta", "message_stop") and (acc["prompt_tokens"] or acc["completion_tokens"]):
        await queue.put({"type": "usage", "usage": {
            "prompt_tokens": acc["prompt_tokens"],
            "completion_tokens": acc["completion_tokens"],
            "total_tokens": acc["prompt_tokens"] + acc["completion_tokens"],
        }})


@app.post("/api/test")
async def api_test(request: Request):
    """「测试连接」按钮：拿这组参数真发一次最小的请求，通不通立刻知道。

    body: {protocol: "openai"|"anthropic", base_url, api_key, model}
    key/地址留空就用已经存好的那份。
    """
    try:
        body = await request.json()
    except Exception:
        body = {}
    protocol = "anthropic" if str(body.get("protocol") or "").lower() == "anthropic" else "openai"
    ag, ch = CONFIG["agent"], CONFIG["chat"]
    # 填了就用填的；没填就退回已经存好的那份（Anthropic 那边没配就用聊天那把 Key —— 同一个账号通常通用）
    base = str(body.get("base_url") or "").strip().rstrip("/") or \
        ((ag["base_url"] or _guess_anthropic_base(ch["base_url"]))
         if protocol == "anthropic" else ch["base_url"])
    key = str(body.get("api_key") or "").strip() or \
        ((ag["api_key"] or ch["api_key"]) if protocol == "anthropic" else ch["api_key"])
    model = str(body.get("model") or "").strip() or \
        ((ag["model"] or ch["model"]) if protocol == "anthropic" else ch["model"]) or "deepseek-flash"
    if not base:
        return {"ok": False, "message": "还没填接口地址"}
    if not key:
        return {"ok": False, "message": "还没填 Key（干活模式也要一把 Anthropic 兼容的 Key）"}

    t0 = time.time()
    try:
        if protocol == "anthropic":
            headers = {"x-api-key": key, "Authorization": f"Bearer {key}",
                       "anthropic-version": "2023-06-01", "content-type": "application/json"}
            r = await app.state.client.post(
                _anthropic_url(base), headers=headers, timeout=60.0,
                json={"model": model, "max_tokens": 64,
                      "messages": [{"role": "user", "content": "只回答两个字：可以"}]})
        else:
            headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
            r = await app.state.client.post(
                base + "/chat/completions", headers=headers, timeout=60.0,
                json={"model": model, "max_tokens": 64, "stream": False,
                      "messages": [{"role": "user", "content": "只回答两个字：可以"}]})
        ms = int((time.time() - t0) * 1000)
        if r.status_code != 200:
            _, msg, hint = humanize_error(r.status_code, r.text)
            return {"ok": False, "message": f"接口拒绝了这个 Key/模型：{msg}", "hint": hint, "ms": ms}
        d = r.json()
        if protocol == "anthropic":
            txt = "".join(x.get("text") or "" for x in (d.get("content") or []) if isinstance(x, dict))
        else:
            txt = ((d.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
        return {"ok": True, "ms": ms, "model": d.get("model") or model,
                "message": f"通了！{ms}ms · 模型 {d.get('model') or model} 回了「{txt.strip()[:20]}」"}
    except Exception as e:
        _, msg, hint = network_error(e)
        return {"ok": False, "message": msg, "hint": hint}


@app.get("/api/balance")
async def api_balance(fresh: int = 0):
    """代理查余额。默认缓存 60 秒,加 ?fresh=1 强制刷新。"""
    cache = app.state.balance_cache
    if not fresh and cache["data"] and time.time() - cache["ts"] < 60:
        return cache["data"]

    if not _have_key:
        return {"ok": False, "message": "没有配置 API Key,查不了余额。"}

    url = CONFIG["base_url"].rstrip("/") + "/user/balance"
    headers = {"Authorization": f"Bearer {CONFIG['api_key']}"}
    try:
        r = await app.state.client.get(url, headers=headers, timeout=20.0)
        if r.status_code != 200:
            code, msg, hint = humanize_error(r.status_code, r.text)
            return {"ok": False, "message": msg, "hint": hint}
        d = r.json()
        infos = d.get("balance_infos") or []
        info = infos[0] if infos else {}
        out = {
            "ok": True,
            "is_available": bool(d.get("is_available")),
            "currency": info.get("currency", "CNY"),
            "total": info.get("total_balance", "?"),
            "topped_up": info.get("topped_up_balance", "?"),
            "ts": time.time(),
        }
        cache["data"], cache["ts"] = out, time.time()
        return out
    except Exception as e:
        _, msg, hint = network_error(e)
        return {"ok": False, "message": msg, "hint": hint}


# ---------------------------- 图片 ----------------------------
@app.post("/api/upload")
async def api_upload(request: Request, ext: str = "png"):
    """前端把压缩过的图片以原始字节 POST 过来,返回一个 id。图片存 data/images/。"""
    raw = await request.body()
    if not raw:
        return JSONResponse({"ok": False, "message": "空文件"}, status_code=400)
    if len(raw) > 12 * 1024 * 1024:
        return JSONResponse({"ok": False, "message": "图片太大了(超过 12MB)"}, status_code=413)
    ext = (ext or "png").lower().lstrip(".")
    if ext not in ("png", "jpg", "jpeg", "webp", "gif"):
        ext = "png"
    img_id = f"{int(time.time() * 1000):x}-{os.urandom(4).hex()}.{ext}"
    (IMG_DIR / img_id).write_bytes(raw)
    return {"ok": True, "id": img_id, "url": f"/api/image/{img_id}", "bytes": len(raw)}


@app.get("/api/image/{img_id}")
async def api_image(img_id: str):
    if not SAFE_IMG.match(img_id):
        return JSONResponse({"ok": False, "message": "非法文件名"}, status_code=400)
    f = (IMG_DIR / img_id).resolve()
    # 双保险：解析出来的路径必须还在 images 文件夹里
    if IMG_DIR.resolve() not in f.parents:
        return JSONResponse({"ok": False, "message": "非法路径"}, status_code=400)
    if not f.exists():
        return JSONResponse({"ok": False, "message": "图片不存在"}, status_code=404)
    return FileResponse(str(f))


def _build_upstream_messages(messages: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """把前端存的对话格式,转成 DeepSeek 要的格式。
       图片 -> base64；附件(文档/音频转写) -> 包成 <file> 文本块塞进消息。"""
    out: List[Dict[str, Any]] = []
    for m in messages:
        role = m.get("role")
        if role not in ("user", "assistant"):
            continue
        text = (m.get("content") or "").strip()
        imgs = [i for i in (m.get("images") or []) if isinstance(i, str)]
        atts = [a for a in (m.get("attach") or []) if isinstance(a, str)]

        if role == "assistant":
            # 助手消息只回传文字;思考过程(reasoning_content)不往回传。
            # 被停止/失败的空气泡直接丢掉,免得 API 报错。
            if text:
                out.append({"role": "assistant", "content": text})
            continue

        # ---- 附件：先在本地解析好的内容包成文本，超出上限就截断 ----
        file_blocks: List[str] = []
        budget = MAX_INJECT_CHARS
        for aid in atts:
            meta = _load_attachment(aid)
            if not meta:
                continue
            body = meta.get("text") or ""
            if body:
                use = body[:budget]
                budget -= len(use)
                cut = "" if len(use) == len(body) else f"\n…（原 {len(body)} 字，按上限截断）"
                file_blocks.append(f'<file name="{meta.get("name")}" type="{meta.get("kind")}">\n{use}{cut}\n</file>')
            # 附件里带的图片（扫描件页/视频帧）也一起发
            for im in (meta.get("images") or [])[:MAX_ATTACH_IMAGES]:
                if len(imgs) < MAX_ATTACH_IMAGES:
                    imgs.append(im)
            if budget <= 0:
                file_blocks.append("…（附件太多，后面的内容没有发送）")
                break
        if file_blocks:
            text = ("\n\n".join(file_blocks) + ("\n\n" + text if text else "")).strip()

        parts: List[Dict[str, Any]] = []
        if imgs:
            for iid in imgs[:MAX_ATTACH_IMAGES]:
                if not SAFE_IMG.match(iid):
                    print(f"[chat] 跳过不合法的图片 id: {iid!r}")
                    continue
                f = IMG_DIR / iid
                if not f.exists():
                    print(f"[chat] 跳过找不到的图片: {iid!r}")
                    continue
                import base64
                b64 = base64.b64encode(f.read_bytes()).decode()
                mime = "image/jpeg" if iid.lower().endswith(("jpg", "jpeg")) else \
                       "image/webp" if iid.lower().endswith("webp") else \
                       "image/gif" if iid.lower().endswith("gif") else "image/png"
                parts.append({"type": "image_url",
                              "image_url": {"url": f"data:{mime};base64,{b64}"}})
        if text:
            parts.append({"type": "text", "text": text})
        if not parts:
            continue
        # 纯文字就发纯字符串(更省 token、兼容性最好)
        out.append({"role": "user", "content": parts[0]["text"] if len(parts) == 1 and not imgs else parts})
    return out


@app.post("/api/chat")
async def api_chat(request: Request):
    """核心:流式转发。浏览器断开 -> 我们也会掐断对 DeepSeek 的请求。"""
    if not _have_key:
        async def _nokey():
            yield 'data: ' + json.dumps({
                "type": "error", "code": "nokey",
                "message": "没有找到 API Key,发不出去。",
                "hint": "在 config.json 里填 api_key,或设置环境变量 DEEPSEEK_API_KEY,"
                        "或确认 ~/.claude/settings.json 里有 ANTHROPIC_AUTH_TOKEN。然后重启本程序。",
            }, ensure_ascii=False) + "\n\n"
            yield 'data: {"type":"done"}\n\n'
        return StreamingResponse(_nokey(), media_type="text/event-stream")

    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)

    messages = body.get("messages") or []
    up_msgs = _build_upstream_messages(messages)
    if not up_msgs:
        async def _empty():
            yield 'data: ' + json.dumps({
                "type": "error", "code": "empty",
                "message": "没有可发送的内容。",
                "hint": "先输入点什么,或贴一张图片。",
            }, ensure_ascii=False) + "\n\n"
            yield 'data: {"type":"done"}\n\n'
        return StreamingResponse(_empty(), media_type="text/event-stream")

    # 模型名不再卡死在内置清单里（中转站的模型名千奇百怪），只要不是空/超长就放行
    model = str(body.get("model") or CONFIG["model"]).strip() or CONFIG["model"]
    if len(model) > 120:
        model = CONFIG["model"]
    try:
        temperature = float(body.get("temperature", CONFIG["temperature"]))
        temperature = max(0.0, min(2.0, temperature))
    except Exception:
        temperature = 0.7
    try:
        max_tokens = int(body.get("max_tokens", CONFIG["max_tokens"]))
        max_tokens = max(1, min(8192, max_tokens))
    except Exception:
        max_tokens = 4096

    system = (body.get("system") or "").strip()
    # 记忆注入：CHAT 模式本来"不认识"用户，把记忆索引塞进 system 就接上了
    if body.get("memory"):
        mem = _memory_text()
        if mem:
            system = (system + "\n\n" if system else "") + (
                "【关于这个用户的长期记忆（来自他本机的记忆库，请直接当作已知事实使用，"
                "不要复述这段、也不要说你读了记忆文件）】\n" + mem
            )

    # 两种协议：OpenAI 兼容（DeepSeek 官方那种）和 Anthropic 兼容（中转站 / Claude 那套）
    anthropic = CONFIG["chat"]["protocol"] == "anthropic"
    if anthropic:
        payload: Dict[str, Any] = {
            "model": model, "stream": True, "temperature": temperature,
            "max_tokens": max_tokens, "messages": _to_anthropic_messages(up_msgs),
        }
        if system:
            payload["system"] = system        # Anthropic 的 system 是顶层参数，不是一条消息
        url = _anthropic_url(CONFIG["chat"]["base_url"])
        headers = {
            # 两种鉴权头都带上：x-api-key 是 Anthropic 原生，Bearer 是中转站常见写法
            "x-api-key": CONFIG["chat"]["api_key"],
            "Authorization": f"Bearer {CONFIG['chat']['api_key']}",
            "anthropic-version": "2023-06-01",
            "Content-Type": "application/json",
            "Accept": "text/event-stream",
        }
    else:
        payload = {
            "model": model, "stream": True, "temperature": temperature,
            "max_tokens": max_tokens, "messages": up_msgs,
        }
        if system:
            payload["messages"] = [{"role": "system", "content": system}] + up_msgs
        url = CONFIG["chat"]["base_url"].rstrip("/") + "/chat/completions"
        headers = {
            "Authorization": f"Bearer {CONFIG['chat']['api_key']}",
            "Content-Type": "application/json",
            "Accept": "text/event-stream",
        }

    # 统计真正发出去的图片数（含附件带来的扫描件页/视频帧），别只看 m["images"]
    n_img = 0
    for m in up_msgs:
        c = m.get("content")
        if isinstance(c, list):
            n_img += sum(1 for p in c if isinstance(p, dict) and p.get("type") == "image_url")
    print(f"[chat] protocol={'anthropic' if anthropic else 'openai'} model={model} "
          f"msgs={len(up_msgs)} imgs={n_img} temp={temperature} max_tokens={max_tokens}")

    client: httpx.AsyncClient = app.state.client
    queue: asyncio.Queue = asyncio.Queue(maxsize=512)

    async def producer():
        """后台任务:读 DeepSeek 的流,塞进队列。"""
        try:
            async with client.stream("POST", url, headers=headers, json=payload) as resp:
                if resp.status_code != 200:
                    raw = (await resp.aread()).decode("utf-8", "replace")
                    code, msg, hint = humanize_error(resp.status_code, raw)
                    print(f"[chat] HTTP {resp.status_code} -> {code}")
                    await queue.put({"type": "error", "code": code,
                                     "message": msg, "hint": hint,
                                     "status": resp.status_code})
                    return
                acc: Dict[str, Any] = {"prompt_tokens": 0, "completion_tokens": 0}
                async for line in resp.aiter_lines():
                    if not line or not line.startswith("data:"):
                        continue
                    data = line[5:].strip()
                    if data == "[DONE]":
                        break
                    try:
                        j = json.loads(data)
                    except Exception:
                        continue
                    if anthropic:
                        await _pump_anthropic(j, queue, acc)
                        continue
                    ch = (j.get("choices") or [{}])[0]
                    delta = ch.get("delta") or {}
                    rc = delta.get("reasoning_content")
                    ct = delta.get("content")
                    if rc:
                        await queue.put({"type": "reasoning", "text": rc})
                    if ct:
                        await queue.put({"type": "content", "text": ct})
                    if j.get("usage"):
                        await queue.put({"type": "usage", "usage": j["usage"]})
                    if ch.get("finish_reason"):
                        await queue.put({"type": "finish", "reason": ch["finish_reason"]})
        except asyncio.CancelledError:
            raise
        except Exception as e:
            code, msg, hint = network_error(e)
            print(f"[chat] 异常 {type(e).__name__}: {e}")
            with contextlib.suppress(Exception):
                await queue.put({"type": "error", "code": code, "message": msg, "hint": hint})
        finally:
            with contextlib.suppress(Exception):
                queue.put_nowait(None)   # 结束哨兵

    async def event_stream():
        task = asyncio.create_task(producer())
        try:
            while True:
                try:
                    item = await asyncio.wait_for(queue.get(), timeout=15.0)
                except asyncio.TimeoutError:
                    yield ": keep-alive\n\n"     # 防止长时间没数据被中间层掐断
                    continue
                if item is None:
                    break
                yield "data: " + json.dumps(item, ensure_ascii=False) + "\n\n"
                if await request.is_disconnected():
                    break
            yield 'data: {"type":"done"}\n\n'
        finally:
            if not task.done():
                task.cancel()
            with contextlib.suppress(asyncio.CancelledError, Exception):
                await task

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


# ==========================================================================
#  干活模式：把网页接到 Claude Code 上
#  原理：调 `claude -p --output-format stream-json`，把它的 JSON 事件流
#        翻译成我们自己的 SSE 推给浏览器。session_id 用 --resume 续聊。
# ==========================================================================
CLAUDE_BIN: Optional[str] = None
for _cand in (
    os.environ.get("CLAUDE_BIN"),
    # 打包版自带的 claude（原生 exe，不需要 Node、不需要 PATH 里有东西）——优先用它
    str(ROOT / "claude" / "claude.exe"),
    str(ROOT / "claude.exe"),
    shutil.which("claude"),
    str(Path.home() / "AppData" / "Roaming" / "npm" / "claude.cmd"),
    str(Path.home() / ".local" / "bin" / "claude"),
):
    if _cand and Path(_cand).exists():
        CLAUDE_BIN = _cand
        break


def _agent_env() -> Dict[str, str]:
    """给 claude 子进程准备环境变量。

    · config.json 里 agent 那一段**完全没配** → 什么都不动，让 claude 读它自己的
      ~/.claude/settings.json（本机一直是这种，能用就别去打扰它）
    · 配了地址/模型但没配单独的 Key → 用聊天那把 Key（同一个账号通常通用；
      ⚠️ 少了这一步，打包给别人的机器上干活模式会因为没 Key 直接失败——踩过）
    """
    env = dict(os.environ)
    ag = CONFIG["agent"]
    if not (ag["base_url"] or ag["api_key"] or ag["model"]):
        return env
    if ag["base_url"]:
        env["ANTHROPIC_BASE_URL"] = ag["base_url"]
    key = ag["api_key"] or CONFIG["chat"]["api_key"]
    if key:
        env["ANTHROPIC_AUTH_TOKEN"] = key
    if ag["model"]:
        env["ANTHROPIC_MODEL"] = ag["model"]
        # 它内部还会去要"小模型"（起标题/压缩上下文），不映射的话会去找一个中转站没有的模型名
        for k in ("ANTHROPIC_DEFAULT_HAIKU_MODEL", "ANTHROPIC_DEFAULT_SONNET_MODEL",
                  "ANTHROPIC_DEFAULT_OPUS_MODEL", "CLAUDE_CODE_SUBAGENT_MODEL",
                  "ANTHROPIC_DEFAULT_HAIKU_MODEL_NAME", "ANTHROPIC_DEFAULT_SONNET_MODEL_NAME",
                  "ANTHROPIC_DEFAULT_OPUS_MODEL_NAME"):
            env.setdefault(k, ag["model"])
    env["DISABLE_AUTOUPDATER"] = "1"     # 自带的那份不让它自己去更新（更新会失败刷屏）
    return env

# 全自动模式：显式放行这一套。
# ⚠️ 实测教训：Windows 上执行命令的工具名是 PowerShell（不是 Bash），
#    而用户 settings.json 里只放行了 Bash(*)，不显式放行的话它连命令都跑不了。
FULL_TOOLS = ("Read,Write,Edit,MultiEdit,Glob,Grep,WebSearch,WebFetch,"
              "PowerShell,Bash,Task,TodoWrite,NotebookEdit,Skill")
# 只读模式：必须用 disallowedTools 明确禁止。
# ⚠️ 实测教训：--allowedTools 是"额外允许"，挡不住默认配置里已放行的 Write(*)，
#    只有 --disallowedTools 才真的拦得住。
DENY_TOOLS = "Write,Edit,MultiEdit,NotebookEdit,Bash,PowerShell,Task,KillShell"


def _translate_agent_event(ev: Dict[str, Any]) -> List[Dict[str, Any]]:
    """把 claude 的一条 JSON 事件翻译成前端要的事件（一条可能翻译出多个）"""
    out: List[Dict[str, Any]] = []
    t = ev.get("type")
    sub = ev.get("subtype")

    if t == "system":
        if sub == "init":
            out.append({"type": "session", "session_id": ev.get("session_id"),
                        "cwd": ev.get("cwd"), "model": ev.get("model")})
        elif sub == "permission_denied":
            out.append({"type": "denied", "name": ev.get("tool_name") or ev.get("tool"),
                        "message": ev.get("message") or "这个工具没被允许"})
        # thinking_tokens 之类的高频噪音直接丢掉
        return out

    if t == "assistant":
        for c in (ev.get("message") or {}).get("content") or []:
            ct = c.get("type")
            if ct == "text" and c.get("text"):
                out.append({"type": "text", "text": c["text"]})
            elif ct == "thinking" and c.get("thinking"):
                out.append({"type": "thinking", "text": c["thinking"]})
            elif ct == "tool_use":
                out.append({"type": "tool", "id": c.get("id"), "name": c.get("name"),
                            "input": c.get("input") or {}})
        return out

    if t == "user":
        for c in (ev.get("message") or {}).get("content") or []:
            if c.get("type") == "tool_result":
                cont = c.get("content")
                if isinstance(cont, list):
                    parts = []
                    for x in cont:
                        if isinstance(x, dict) and x.get("type") == "text":
                            parts.append(x.get("text") or "")
                        else:
                            parts.append(str(x))
                    cont = "\n".join(parts)
                out.append({"type": "tool_result", "id": c.get("tool_use_id"),
                            "is_error": bool(c.get("is_error")),
                            "content": (cont or "")[:20000]})
        return out

    if t == "result":
        u = ev.get("usage") or {}
        out.append({"type": "agent_usage",
                    "input_tokens": u.get("input_tokens"),
                    "output_tokens": u.get("output_tokens"),
                    "cost_usd": ev.get("total_cost_usd"),
                    "duration_ms": ev.get("duration_ms"),
                    "is_error": bool(ev.get("is_error")),
                    "result": (ev.get("result") or "")[:4000],
                    "stop_reason": ev.get("stop_reason")})
        return out

    return out


def _kill_tree(pid: int) -> None:
    """连子进程一起杀掉（Windows 上必须 /T，否则 node 会赖着）"""
    with contextlib.suppress(Exception):
        if sys.platform.startswith("win"):
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(pid)],
                           capture_output=True, creationflags=0x08000000)
        else:
            os.killpg(os.getpgid(pid), signal.SIGKILL)


@app.get("/api/agent/available")
async def agent_available():
    return {"ok": bool(CLAUDE_BIN), "bin": CLAUDE_BIN,
            "hint": "" if CLAUDE_BIN else
            "找不到 claude 程序：装一下（npm install -g @anthropic-ai/claude-code），"
            "或把 claude.exe 放到程序目录的 claude\\ 文件夹里"}


@app.post("/api/agent")
async def api_agent(request: Request):
    """干活模式：网页 → claude CLI → 事件流回网页"""
    if not CLAUDE_BIN:
        async def _nobin():
            yield 'data: ' + json.dumps({
                "type": "error", "code": "nobinary",
                "message": "找不到 claude 命令，没法干活。",
                "hint": "装一下：npm install -g @anthropic-ai/claude-code",
            }, ensure_ascii=False) + "\n\n"
            yield 'data: {"type":"done"}\n\n'
        return StreamingResponse(_nobin(), media_type="text/event-stream")

    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)

    prompt = (body.get("prompt") or "").strip()
    if not prompt:
        return JSONResponse({"ok": False, "message": "没说要干什么"}, status_code=400)

    # 附件也要给 Claude Code 看到！否则「贴了图但 agent 看不见」——
    # 图片/扫描件给本地路径让它自己 Read；已经解析成文字的（文档/录音）直接内联。
    atts = [a for a in (body.get("attach") or []) if isinstance(a, str)]
    if atts:
        blocks: List[str] = []
        budget = MAX_INJECT_CHARS
        for aid in atts:
            meta = _load_attachment(aid)
            if not meta:
                continue
            name = meta.get("name") or aid
            for im in (meta.get("images") or []):
                f = IMG_DIR / im
                if f.exists():
                    blocks.append(f"- 「{name}」的本地图片路径：{f}\n  （请用 Read 工具打开它，看清楚内容再回答）")
            if meta.get("text") and meta.get("kind") != "image":
                use = meta["text"][:budget]
                budget -= len(use)
                cut = "" if len(use) == len(meta["text"]) else "\n…（按上限截断）"
                blocks.append(f'<file name="{name}">\n{use}{cut}\n</file>')
        if blocks:
            print(f"[agent] 附带 {len(blocks)} 份材料一起交给它")
            prompt = ("【用户附带的材料】\n" + "\n".join(blocks)
                      + "\n\n【用户说的话】\n" + prompt)

    # 起始目录只决定 claude 从哪儿起步，**不是沙箱**（它照样能访问别的盘/目录）。
    # 所以填错了不值得报错打断用户——直接用家目录兜底，日志里记一笔就行。
    workspace = (str(body.get("workspace") or "").strip()
                 or CONFIG["agent"]["workspace"]
                 or str(Path.home()))
    if not Path(workspace).is_dir():
        # 配好了起始目录但还没建出来 → 帮它建一个（别悄悄换回家目录：
        # 实测踩过——目录不存在就退回家目录，结果文件全建到 C:\Users\你\ 底下去了）
        with contextlib.suppress(Exception):
            Path(workspace).mkdir(parents=True, exist_ok=True)
    if not Path(workspace).is_dir():
        print(f"[agent] 起始目录不存在也建不出来({workspace})，改用家目录 {Path.home()}")
        workspace = str(Path.home())
    # 这个目录也共用同一份记忆（claude 按 cwd 找记忆，没见过的新目录默认空白——踩过）
    _note = _link_memory_to_shared(CLAUDE_PROJECTS / _proj_slug(workspace))
    if _note:
        print(f"[agent] {_note}")

    session_id = (body.get("session_id") or "").strip()
    readonly = bool(body.get("readonly"))

    WEB_CONTEXT_NOTE = (
        "你正在一个本地的网页工作台里被调用（用户在看浏览器，不是终端 TUI）。"
        "有几件事你需要知道：① 用户无法做交互式确认，需要他决策就在回答里写清楚，别指望弹窗；"
        "② 你的会话记录和长期记忆跟终端里是同一套，不用重新认识用户；"
        "③ 尽量用简洁的结论收尾，他会直接在网页上看。"
    )
    cmd = [CLAUDE_BIN, "-p", "--output-format", "stream-json", "--verbose",
           "--append-system-prompt", WEB_CONTEXT_NOTE]
    if session_id:
        cmd += ["--resume", session_id]
    if readonly:
        cmd += ["--disallowedTools", DENY_TOOLS]      # 只读：只许看，不许改
    else:
        cmd += ["--allowedTools", FULL_TOOLS]         # 全自动：这一套都放行
        # 2026-09-13 补丁：网页端弹不出权限确认（不询问=直接拒），工作区之外的改动（如共享记忆目录）会被静默拒
        cmd += ["--permission-mode", "bypassPermissions"]
        try:
            if MEM_DIR.is_dir():
                cmd += ["--add-dir", str(MEM_DIR)]
        except Exception:
            pass

    print(f"[agent] cwd={workspace} resume={'yes' if session_id else 'no'} "
          f"readonly={readonly} prompt={prompt[:60]!r}")

    CREATE_NO_WINDOW = 0x08000000
    try:
        proc = subprocess.Popen(
            cmd, cwd=workspace, env=_agent_env(),
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, encoding="utf-8", errors="replace", bufsize=1,
            creationflags=CREATE_NO_WINDOW if sys.platform.startswith("win") else 0,
        )
    except Exception as e:
        async def _fail():
            yield 'data: ' + json.dumps({
                "type": "error", "code": "spawn", "message": f"起不来 claude 进程：{e}",
                "hint": f"命令是 {cmd[0]}，确认它能用。"}, ensure_ascii=False) + "\n\n"
            yield 'data: {"type":"done"}\n\n'
        return StreamingResponse(_fail(), media_type="text/event-stream")

    # 用管道把问题喂进去（避免命令行引号问题）
    with contextlib.suppress(Exception):
        proc.stdin.write(prompt)
        proc.stdin.close()

    # 登记这次任务：网页断开（刷新/关标签）不会杀它，跑完的结果留 2 小时，
    # 前端凭 run_id 取回（见 /api/agent/result）。
    _prune_agent_runs()
    run_id = secrets.token_hex(6)
    AGENT_RUNS[run_id] = {"proc": proc, "events": [], "text": "", "done": False,
                          "sid": session_id or "", "ts": time.time()}

    loop = asyncio.get_running_loop()
    queue: asyncio.Queue = asyncio.Queue(maxsize=2000)
    err_lines: List[str] = []

    def _pump_stdout():
        try:
            for line in proc.stdout:
                line = line.strip()
                if not line:
                    continue
                try:
                    ev = json.loads(line)
                except Exception:
                    continue
                for item in _translate_agent_event(ev):
                    _remember_run_event(run_id, item)      # 存一份，刷新后能取回
                    loop.call_soon_threadsafe(queue.put_nowait, item)
        except Exception as e:
            loop.call_soon_threadsafe(queue.put_nowait,
                                      {"type": "error", "message": f"读输出出错：{e}"})

    def _pump_stderr():
        with contextlib.suppress(Exception):
            for line in proc.stderr:
                err_lines.append(line.rstrip())
                del err_lines[:-40]

    th_out = threading.Thread(target=_pump_stdout, daemon=True)
    th_err = threading.Thread(target=_pump_stderr, daemon=True)
    th_out.start()
    th_err.start()

    def _wait_done():
        rc = proc.wait()
        run = AGENT_RUNS.get(run_id)
        if run is not None:
            run["done"] = True
            run["rc"] = rc
        loop.call_soon_threadsafe(queue.put_nowait, {"type": "exit", "code": rc})
        if rc != 0:
            tail = "\n".join(err_lines[-12:])
            loop.call_soon_threadsafe(queue.put_nowait, {
                "type": "error", "code": "exit",
                "message": f"claude 退出了（代码 {rc}）",
                "hint": tail[-1200:] or "没有更多信息，看看是不是工作目录不对、或者 claude 没登录。",
            })
        loop.call_soon_threadsafe(queue.put_nowait, None)

    threading.Thread(target=_wait_done, daemon=True).start()

    async def event_stream():
        # 第一条就把 run_id 告诉前端：刷新回来后凭它把结果取回来／主动停止
        yield 'data: ' + json.dumps({"type": "run", "run_id": run_id},
                                    ensure_ascii=False) + "\n\n"
        try:
            while True:
                try:
                    item = await asyncio.wait_for(queue.get(), timeout=20.0)
                except asyncio.TimeoutError:
                    yield ": keep-alive\n\n"
                    if await request.is_disconnected():
                        break
                    continue
                if item is None:
                    break
                yield "data: " + json.dumps(item, ensure_ascii=False) + "\n\n"
                if await request.is_disconnected():
                    break
            yield 'data: {"type":"done"}\n\n'
        finally:
            # 网页断开（刷新 / 关标签 / 断网）**不杀任务**：它在后台继续跑，结果留在
            # AGENT_RUNS 里，前端刷新回来凭 run_id 取回。
            # 真正要停 = 用户点停止按钮 → 前端先调 /api/agent/stop，那才杀进程。
            if proc.poll() is None:
                print(f"[agent] 网页断开了，任务继续在后台跑（run={run_id}，要停请点停止按钮）")

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache, no-transform",
                 "Connection": "keep-alive", "X-Accel-Buffering": "no"},
    )


# ==========================================================================
#  在跑的 agent 任务：刷新页面 ≠ 中断
#  run_id 由 /api/agent 的第一条事件下发；跑完的结果在这里留 2 小时，
#  前端刷新回来凭 run_id 把最终回答填回气泡（也可以直接追问，会话没断）。
# ==========================================================================
AGENT_RUNS: Dict[str, Dict[str, Any]] = {}
AGENT_RUN_TTL = 2 * 3600


def _remember_run_event(run_id: str, item: Dict[str, Any]) -> None:
    """每个事件在内存里留一份（最近 4000 条 + 最终回答）"""
    run = AGENT_RUNS.get(run_id)
    if run is None:
        return
    run["events"].append(item)
    if len(run["events"]) > 4000:
        del run["events"][:-4000]
    t = item.get("type")
    if t == "session":
        run["sid"] = item.get("session_id") or run["sid"]
    elif t == "text":
        run["text"] = (run["text"] + (item.get("text") or ""))[-40000:]
    elif t == "agent_usage":
        run["usage"] = {"input_tokens": item.get("input_tokens"),
                        "output_tokens": item.get("output_tokens")}
        if item.get("result"):
            run["text"] = item["result"]


def _prune_agent_runs() -> None:
    now = time.time()
    for k in [k for k, v in AGENT_RUNS.items() if now - v.get("ts", 0) > AGENT_RUN_TTL]:
        AGENT_RUNS.pop(k, None)


@app.get("/api/agent/result")
async def api_agent_result(run_id: str):
    """凭 run_id 看一个 agent 任务：还在跑吗？跑完没有？最终回答是什么？"""
    run = AGENT_RUNS.get(run_id)
    if not run:
        return {"ok": False, "message": "没有这个任务（可能跑完超过 2 小时，或服务重启过）"}
    return {"ok": True, "done": bool(run.get("done")), "text": run.get("text") or "",
            "session_id": run.get("sid") or "", "usage": run.get("usage"),
            "killed": bool(run.get("killed")), "age": int(time.time() - run.get("ts", 0))}


@app.post("/api/agent/stop")
async def api_agent_stop(request: Request):
    """真正的中断：前端点停止按钮时调这里（刷新页面**不**算停止）"""
    body: Dict[str, Any] = {}
    with contextlib.suppress(Exception):
        body = await request.json()
    rid = str(body.get("run_id") or "")
    run = AGENT_RUNS.get(rid)
    if not run:
        return {"ok": False, "message": "没找到这个任务（可能已经跑完了）"}
    proc = run.get("proc")
    if proc is not None and proc.poll() is None:
        await asyncio.to_thread(_kill_tree, proc.pid)
        run["killed"] = True
        print(f"[agent] 用户点了停止，已杀掉 run={rid}")
    return {"ok": True}


# ==========================================================================
#  更新干活模式的「大脑」：把自带的 claude.exe 换成最新版
#
#  为什么不用 claude 自带的 `claude update`：它去 downloads.claude.ai 拿版本，
#  国内直连被拒（2026-09-13 实测 ECONNREFUSED，重试 3 次失败）。
#  这里改走国内 npm 镜像（registry.npmmirror.com → cdn.npmmirror.com，
#  实测 8 MB/s，98MB 的包十几秒下完），下完当场验版本、验能跑，再替换。
# ==========================================================================
NPMMIRROR = "https://registry.npmmirror.com"
CLAUDE_PKG = "@anthropic-ai/claude-code"
# 原生二进制是单独的"平台包"（npm 主包只有 0.2MB，安装时再去拉对应平台的那个）
CLAUDE_PLATFORM_PKG = {
    "win32": {"x86_64": "claude-code-win32-x64", "arm64": "claude-code-win32-arm64"},
    "darwin": {"x86_64": "claude-code-darwin-x64", "arm64": "claude-code-darwin-arm64"},
    "linux": {"x86_64": "claude-code-linux-x64", "arm64": "claude-code-linux-arm64"},
}


def _platform_pkg() -> str:
    import platform as _pf
    mach = {"AMD64": "x86_64", "x86_64": "x86_64", "ARM64": "arm64", "aarch64": "arm64"}.get(_pf.machine(), "")
    return CLAUDE_PLATFORM_PKG.get(sys.platform.split("-")[0], {}).get(mach, "")


def _ver_tuple(v: str) -> tuple:
    return tuple(int(x) for x in re.findall(r"\d+", v or "")[:4]) or (0,)


def _claude_version_of(exe: str) -> str:
    """跑一下 `exe --version` 拿 x.y.z（拿不到就返回空串）"""
    try:
        r = subprocess.run([exe, "--version"], capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=90,
                           creationflags=0x08000000 if sys.platform.startswith("win") else 0)
        m = re.search(r"(\d+\.\d+\.\d+)", (r.stdout or "") + (r.stderr or ""))
        return m.group(1) if m else ""
    except Exception:
        return ""


def _claude_busy() -> bool:
    """自带那个 claude.exe 是不是正在干活（正在跑就替换不了文件）"""
    if not CLAUDE_BIN or not sys.platform.startswith("win"):
        return False
    esc = str(CLAUDE_BIN).replace("'", "''")
    ps = ("Get-Process claude -ErrorAction SilentlyContinue | "
          f"Where-Object {{ $_.Path -eq '{esc}' }} | Measure-Object | "
          "Select-Object -ExpandProperty Count")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, text=True, timeout=40,
                           creationflags=0x08000000)
        return int((r.stdout or "0").strip() or 0) > 0
    except Exception:
        return True          # 查不出来就当作在忙：宁可让她等一下，也别把 exe 换坏


async def _claude_latest() -> str:
    """问国内镜像最新版是多少"""
    r = await app.state.client.get(f"{NPMMIRROR}/{CLAUDE_PKG}/latest", timeout=40.0)
    r.raise_for_status()
    return str((r.json() or {}).get("version") or "")


@app.get("/api/claude-update")
async def api_claude_update_check():
    """查有没有新版 → {current, latest, has_update}"""
    current = await asyncio.to_thread(_claude_version_of, CLAUDE_BIN) if CLAUDE_BIN else ""
    latest, err = "", ""
    try:
        latest = await _claude_latest()
    except Exception as e:
        err = f"查不到最新版（{type(e).__name__}）——断网的话连上再试"
    return {"ok": True, "current": current, "latest": latest, "bin": CLAUDE_BIN,
            "has_update": bool(current and latest and _ver_tuple(latest) > _ver_tuple(current)),
            "message": err}


@app.post("/api/claude-update")
async def api_claude_update_do():
    """下载新版并替换自带的 claude.exe（旧版改名留档 .bak-<版本>，出问题能换回来）"""
    if not CLAUDE_BIN:
        return JSONResponse({"ok": False, "message": "没有找到 claude 程序，没法更新"}, status_code=400)
    if not str(CLAUDE_BIN).lower().endswith(".exe"):
        # npm 装的是个 claude.cmd 外壳，用这种方式替换会把那个安装弄坏 → 让它走 npm
        return JSONResponse({"ok": False, "message":
                             "当前用的是 npm 安装的 claude（claude.cmd），请用 npm 更新："
                             "npm install -g @anthropic-ai/claude-code@latest"}, status_code=400)
    if await asyncio.to_thread(_claude_busy):
        return JSONResponse({"ok": False, "message": "它正在干活，等这次跑完再更新（或先按 Esc 中断）"},
                            status_code=409)
    pkg = _platform_pkg()
    if not pkg:
        return JSONResponse({"ok": False, "message": f"不认识这个平台：{sys.platform}，没法自动更新"}, status_code=400)
    try:
        latest = await _claude_latest()
    except Exception as e:
        return JSONResponse({"ok": False, "message": f"连不上国内镜像：{type(e).__name__}"}, status_code=502)
    if not latest:
        return JSONResponse({"ok": False, "message": "镜像没给出最新版本号"}, status_code=502)
    current = await asyncio.to_thread(_claude_version_of, CLAUDE_BIN)
    if current and _ver_tuple(latest) <= _ver_tuple(current):
        return {"ok": True, "updated": False, "from": current, "to": current,
                "message": f"已经是最新版（{current}）"}

    pkg_full = f"@anthropic-ai/{pkg}"                    # @anthropic-ai/claude-code-win32-x64
    url = f"{NPMMIRROR}/{pkg_full}/-/{pkg}-{latest}.tgz"
    tmpdir = DATA_DIR / "update"
    tmpdir.mkdir(parents=True, exist_ok=True)
    tgz = tmpdir / f"{pkg}-{latest}.tgz"
    print(f"[update] 下载 claude {latest}：{url}")
    try:
        # follow_redirects：npmmirror 会 302 到 cdn.npmmirror.com，不跟就会直接报 302
        async with app.state.client.stream(
                "GET", url, follow_redirects=True,
                timeout=httpx.Timeout(connect=20.0, read=900.0, write=60.0, pool=20.0)) as resp:
            resp.raise_for_status()
            with open(tgz, "wb") as f:
                async for chunk in resp.aiter_bytes(1 << 20):
                    f.write(chunk)
    except Exception as e:
        with contextlib.suppress(Exception):
            tgz.unlink()
        return JSONResponse({"ok": False, "message": f"下载失败：{type(e).__name__} {e}"}, status_code=502)

    import tarfile
    new_exe = tmpdir / ("claude-new.exe" if sys.platform.startswith("win") else "claude-new")
    try:
        with tarfile.open(tgz, "r:gz") as tf:
            member = None
            for m in tf.getmembers():
                if m.name.replace("\\", "/").lower() in (
                        "package/claude.exe", "package/bin/claude.exe",
                        "package/claude", "package/bin/claude"):
                    member = m
                    break
            if member is None:
                return JSONResponse({"ok": False, "message": "下载包里没找到 claude 可执行文件"}, status_code=502)
            src = tf.extractfile(member)
            with open(new_exe, "wb") as f:
                shutil.copyfileobj(src, f)
    except Exception as e:
        return JSONResponse({"ok": False, "message": f"解包失败：{type(e).__name__} {e}"}, status_code=500)

    got = await asyncio.to_thread(_claude_version_of, str(new_exe))
    if got != latest:            # 验版本 + 验它真能跑起来，不合格就不换
        with contextlib.suppress(Exception):
            new_exe.unlink()
        return JSONResponse({"ok": False, "message":
                             f"下下来的文件不对（期望 {latest}，实际 {got or '跑不起来'}），已放弃替换"},
                            status_code=502)

    target = Path(CLAUDE_BIN)
    bak = Path(str(target) + f".bak-{current or 'old'}")
    try:
        if current:
            with contextlib.suppress(Exception):
                os.replace(str(target), str(bak))      # 改名留档，不占额外空间
        os.replace(str(new_exe), str(target))          # 同盘替换，原子操作
    except Exception as e:
        with contextlib.suppress(Exception):           # 尽力还原，别把 claude.exe 弄没了
            if bak.exists() and not target.exists():
                os.replace(str(bak), str(target))
        return JSONResponse({"ok": False, "message":
                             f"替换失败（多半是旧进程还占着文件，稍后再试）：{type(e).__name__} {e}"},
                            status_code=500)

    with contextlib.suppress(Exception):
        tgz.unlink()
    print(f"[update] claude {current} → {latest} 完成，旧版留在 {bak.name}")
    return {"ok": True, "updated": True, "from": current, "to": latest,
            "message": f"更新完成：{current} → {latest}（下一次干活就用新版，旧版备份在 app\\claude\\）"}


# ==========================================================================
#  附件：文档 / 图片 / 音频 / 视频
#  思路：能在本地解析的就在本地解析（不花钱），只有"音频转文字"要调外部 ASR
#        （复用用户已有的硅基流动密钥，实测有效）
# ==========================================================================
FILES_DIR = DATA_DIR / "files"
FILES_DIR.mkdir(parents=True, exist_ok=True)

# 注入给模型的正文上限（字符）。超了截断并告诉用户
MAX_INJECT_CHARS = 40000
MAX_ATTACH_IMAGES = 8          # 一次最多塞几张图（扫描件/视频帧）

TEXT_EXT = {"txt", "md", "markdown", "csv", "tsv", "json", "yaml", "yml", "log",
            "py", "js", "ts", "html", "css", "xml", "ini", "cfg", "toml", "sh",
            "bat", "cmd", "ps1", "c", "h", "cpp", "hpp", "java", "go", "rs", "sql"}
IMAGE_EXT = {"png", "jpg", "jpeg", "webp", "gif", "bmp"}
AUDIO_EXT = {"mp3", "wav", "m4a", "aac", "flac", "ogg", "wma", "amr", "opus"}
VIDEO_EXT = {"mp4", "mkv", "mov", "avi", "webm", "flv", "wmv", "m4v"}


def _kind_of(ext: str) -> str:
    if ext in IMAGE_EXT: return "image"
    if ext == "pdf": return "pdf"
    if ext in ("docx", "doc"): return "doc"
    if ext in ("xlsx", "xls", "csv"): return "sheet"
    if ext in ("pptx", "ppt"): return "slides"
    if ext in AUDIO_EXT: return "audio"
    if ext in VIDEO_EXT: return "video"
    if ext in TEXT_EXT: return "text"
    return "unknown"


def _read_text_file(p: Path, limit: int = 200000) -> str:
    for enc in ("utf-8", "gbk", "utf-16", "latin-1"):
        try:
            return p.read_text(encoding=enc)[:limit]
        except Exception:
            continue
    return ""


def _extract_pdf(p: Path) -> tuple[str, list[str], str]:
    """返回 (正文, 渲染出的图片名列表, 说明)"""
    try:
        import pymupdf as fitz  # PyMuPDF（新版改名叫 pymupdf 了，老写法 fitz 仍然能用）
    except ImportError:
        try:
            import fitz
        except ImportError:
            return "", [], "缺 PyMuPDF：pip install pymupdf"
    doc = fitz.open(str(p))
    pages = doc.page_count
    text = "\n".join((doc[i].get_text() or "") for i in range(pages))
    # 判定"有没有文字层"：扫描件每页几乎抽不出字，文字版哪怕只有个标题也有几十字。
    # 阈值取 20 字/页 —— 宁可误判成扫描件（多花点 token 读图），也别把扫描件当空文本发出去。
    if len(text.strip()) >= 20 * max(pages, 1):
        return text, [], f"{pages} 页，抽出 {len(text)} 字"
    # 扫描件：渲染成图，交给视觉模型读
    imgs, n = [], min(pages, MAX_ATTACH_IMAGES)
    for i in range(n):
        pix = doc[i].get_pixmap(matrix=fitz.Matrix(2, 2))
        fn = f"{int(time.time()*1000):x}-pdf{i}.png"
        pix.save(str(IMG_DIR / fn))
        imgs.append(fn)
    note = f"{pages} 页，正文层为空（扫描件）→ 已渲染前 {n} 页为图片让模型读"
    return "", imgs, note


def _extract_docx(p: Path) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "缺 python-docx：pip install python-docx"
    d = docx.Document(str(p))
    parts = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    txt = "\n".join(parts)
    return txt, f"抽出 {len(txt)} 字（含 {len(d.tables)} 个表格）"


def _extract_xlsx(p: Path, max_rows: int = 300) -> tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "", "缺 openpyxl：pip install openpyxl"
    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"### 工作表：{ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                lines.append(f"...（超过 {max_rows} 行已截断）")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    txt = "\n".join(lines)
    return txt, f"{len(wb.worksheets)} 个工作表，抽出 {len(txt)} 字"


def _extract_pptx(p: Path) -> tuple[str, str]:
    try:
        from pptx import Presentation
    except ImportError:
        return "", "缺 python-pptx：pip install python-pptx"
    prs = Presentation(str(p))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"### 第 {i} 页")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text)
    txt = "\n".join(lines)
    return txt, f"{len(prs.slides)} 页，抽出 {len(txt)} 字"


def _asr_key() -> str:
    k = (CONFIG.get("asr_key") or "").strip()          # 网页设置里填的（打包版走这条）
    if k:
        return k
    k = (os.environ.get("SILICONFLOW_API_KEY") or "").strip()
    if k:
        return k
    for p in (Path.home() / ".vision" / "api_key_siliconflow.txt",):
        with contextlib.suppress(Exception):
            v = p.read_text(encoding="utf-8").strip().splitlines()[0].strip()
            if v:
                return v
    return ""


def _transcribe(p: Path) -> tuple[str, str]:
    """音频转文字：走硅基流动（OpenAI 兼容的 /audio/transcriptions）"""
    key = _asr_key()
    if not key:
        return "", "没有语音转文字密钥：需要硅基流动的 key（~/.vision/api_key_siliconflow.txt）"
    try:
        import requests as _rq
        with open(p, "rb") as f:
            r = _rq.post(
                "https://api.siliconflow.cn/v1/audio/transcriptions",
                headers={"Authorization": f"Bearer {key}"},
                files={"file": (p.name, f, "application/octet-stream")},
                data={"model": "FunAudioLLM/SenseVoiceSmall"},
                timeout=900,
            )
        if r.status_code != 200:
            return "", f"语音服务返回 {r.status_code}：{r.text[:200]}"
        js = r.json()
        t = js.get("text") or ""
        return t, f"转写出 {len(t)} 字"
    except Exception as e:
        return "", f"转写失败：{type(e).__name__} {e}"


_FFMPEG_CACHE: Dict[str, str] = {}


def _ffmpeg() -> str:
    """找 ffmpeg：PATH 里没有就去 winget 装的目录里翻（winget 装的要重开会话才进 PATH）"""
    if _FFMPEG_CACHE.get("ffmpeg") is not None:
        return _FFMPEG_CACHE["ffmpeg"]
    found = shutil.which("ffmpeg") or ""
    if not found:
        pats = [
            # 打包版可以往程序目录的 bin\ 里丢一个 ffmpeg.exe（体积太大没跟着打包）
            str(ROOT / "bin" / "ffmpeg.exe"),
            str(ROOT / "ffmpeg.exe"),
            str(Path.home() / "AppData/Local/Microsoft/WinGet/Packages/Gyan.FFmpeg*/**/bin/ffmpeg.exe"),
            "C:/ffmpeg/bin/ffmpeg.exe",
            "C:/Program Files/ffmpeg/bin/ffmpeg.exe",
        ]
        for pat in pats:
            hits = glob.glob(pat, recursive=True)
            if hits:
                found = hits[0]
                break
    _FFMPEG_CACHE["ffmpeg"] = found
    if found:
        print(f"[ffmpeg] 用这个：{found}")
    return found


def _ffprobe() -> str:
    if _FFMPEG_CACHE.get("ffprobe") is not None:
        return _FFMPEG_CACHE["ffprobe"]
    p = shutil.which("ffprobe") or ""
    if not p:
        ff = _ffmpeg()
        if ff:
            cand = str(Path(ff).with_name("ffprobe.exe" if os.name == "nt" else "ffprobe"))
            p = cand if Path(cand).exists() else ""
    _FFMPEG_CACHE["ffprobe"] = p
    return p


def _extract_video(p: Path) -> tuple[str, list[str], str]:
    """视频：抽 N 帧 + 抽音轨转写"""
    ff = _ffmpeg()
    if not ff:
        return "", [], "没装 ffmpeg，视频处理不了（装一下：winget install Gyan.FFmpeg）"
    imgs, note = [], []
    dur = 0.0
    try:
        # 时长直接从 ffmpeg 的横幅里读（比 ffprobe 可靠，不依赖它单独存在）
        r0 = subprocess.run([ff, "-i", str(p)], capture_output=True, text=True,
                            encoding="utf-8", errors="replace", timeout=60)
        m = re.search(r"Duration:\s*(\d+):(\d+):([\d.]+)", r0.stderr or "")
        if m:
            dur = int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3))
        n = 8
        for i in range(n):
            ts = (dur * (i + 0.5) / n) if dur > 0.5 else i * 0.8
            fn = f"{int(time.time()*1000):x}-v{i}.png"
            subprocess.run([ff, "-ss", f"{ts:.2f}", "-i", str(p), "-frames:v", "1",
                            "-vf", "scale=768:-1", "-y", str(IMG_DIR / fn)],
                           capture_output=True, timeout=120)
            if (IMG_DIR / fn).exists():
                imgs.append(fn)
        note.append(f"抽了 {len(imgs)} 帧" + (f"（时长 {dur:.0f}s）" if dur else ""))
    except Exception as e:
        note.append(f"抽帧失败：{e}")
    # 音轨 → 转写
    try:
        tmp = FILES_DIR / f"_audio_{p.stem[:20]}.mp3"
        subprocess.run([ff, "-i", str(p), "-vn", "-ac", "1", "-ar", "16000", "-y", str(tmp)],
                       capture_output=True, timeout=600)
        if tmp.exists():
            t, tn = _transcribe(tmp)
            with contextlib.suppress(Exception):
                tmp.unlink()
            if t:
                note.append(f"音轨转写 {len(t)} 字")
                return t, imgs, "；".join(note)
    except Exception as e:
        note.append(f"抽音轨失败：{e}")
    return "", imgs, "；".join(note)


def _parse_attachment(p: Path, kind: str) -> Dict[str, Any]:
    """真正干活的地方：返回 {text, images, note}"""
    if kind == "text":
        t = _read_text_file(p)
        return {"text": t, "images": [], "note": f"文本 {len(t)} 字"}
    if kind == "pdf":
        t, imgs, note = _extract_pdf(p)
        return {"text": t, "images": imgs, "note": note}
    if kind == "doc":
        t, note = _extract_docx(p); return {"text": t, "images": [], "note": note}
    if kind == "sheet" and p.suffix.lower() in (".xlsx", ".xls"):
        t, note = _extract_xlsx(p); return {"text": t, "images": [], "note": note}
    if kind == "sheet":
        t = _read_text_file(p); return {"text": t, "images": [], "note": f"CSV {len(t)} 字"}
    if kind == "slides":
        t, note = _extract_pptx(p); return {"text": t, "images": [], "note": note}
    if kind == "audio":
        t, note = _transcribe(p); return {"text": t, "images": [], "note": note}
    if kind == "video":
        t, imgs, note = _extract_video(p); return {"text": t, "images": imgs, "note": note}
    return {"text": "", "images": [], "note": "这个格式暂时认不出来"}


@app.get("/api/health")
async def api_health():
    return {"ok": True, "t": time.time()}


@app.post("/api/ping")
async def api_ping():
    """写操作连通性自检（前端会用它确认"令牌 + Origin 都放行"）"""
    return {"ok": True, "t": time.time()}


# ==========================================================================
#  Git：看改动 / 看 diff / 一键还原 / 提交
#  不自己造快照系统，直接复用 git（怕 AI 改坏代码时的救命按钮）
# ==========================================================================
def _git(args: List[str], cwd: str, timeout: int = 30) -> tuple[int, str]:
    try:
        # core.quotepath=false：不加这个，git 会把中文文件名转义成 \346\210\220 这种八进制
        r = subprocess.run(["git", "-c", "core.quotepath=false"] + args,
                           cwd=cwd, capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=timeout,
                           creationflags=0x08000000 if sys.platform.startswith("win") else 0)
        return r.returncode, (r.stdout or "") + (r.stderr or "")
    except FileNotFoundError:
        return 127, "没装 git"
    except subprocess.TimeoutExpired:
        return 124, "git 命令超时"


def _safe_ws(ws: str) -> Optional[Path]:
    p = Path(ws or "").expanduser()
    return p if p.is_dir() else None


@app.get("/api/git")
async def api_git(ws: str = ""):
    """看一眼这个目录的 git 状态（不是仓库就明确说）"""
    p = _safe_ws(ws)
    if not p:
        return {"ok": False, "is_repo": False, "message": f"目录不存在：{ws}"}
    code, out = await asyncio.to_thread(_git, ["rev-parse", "--is-inside-work-tree"], str(p))
    if code != 0 or "true" not in out:
        return {"ok": True, "is_repo": False, "ws": str(p),
                "message": "这个目录还不是 git 仓库（可以先初始化，之后改动就能一键回滚）"}

    _, branch = await asyncio.to_thread(_git, ["branch", "--show-current"], str(p))
    _, porcelain = await asyncio.to_thread(_git, ["status", "--porcelain=v1"], str(p))
    files = []
    for line in porcelain.splitlines():
        if len(line) < 4:
            continue
        st, path = line[:2].strip() or "??", line[3:].strip().strip('"')
        files.append({"status": st, "path": path})
    # 有没有"存档点"（没有任何提交的仓库不能还原——还原等于把文件全删了）
    code_h, _ = await asyncio.to_thread(_git, ["rev-parse", "--verify", "HEAD"], str(p))
    has_commits = code_h == 0
    log = ""
    if has_commits:
        _, log = await asyncio.to_thread(_git, ["log", "-1", "--format=%h %s"], str(p))
    return {"ok": True, "is_repo": True, "ws": str(p), "branch": branch.strip(),
            "files": files, "head": log.strip(), "has_commits": has_commits}


@app.get("/api/git/diff")
async def api_git_diff(ws: str = "", path: str = ""):
    """单个文件的 diff（已暂存 + 未暂存都算上）"""
    p = _safe_ws(ws)
    if not p or not path:
        return {"ok": False, "message": "参数不对"}
    if ".." in path or path.startswith("/") or ":" in path:
        return {"ok": False, "message": "非法路径"}
    _, untracked = await asyncio.to_thread(_git, ["ls-files", "--others", "--exclude-standard", "--", path], str(p))
    if untracked.strip():
        # 新文件没有 diff，直接把内容给出来
        f = p / path
        with contextlib.suppress(Exception):
            txt = f.read_text(encoding="utf-8", errors="replace")[:20000]
            return {"ok": True, "diff": f"（新文件，共 {f.stat().st_size} 字节）\n\n" + txt, "is_new": True}
        return {"ok": True, "diff": "（新文件）", "is_new": True}
    _, d = await asyncio.to_thread(_git, ["diff", "HEAD", "--", path], str(p))
    if not d.strip():
        _, d = await asyncio.to_thread(_git, ["diff", "--", path], str(p))
    return {"ok": True, "diff": d[:60000]}


@app.post("/api/git/action")
async def api_git_action(request: Request):
    """写操作：init / restore（还原单个或全部）/ commit"""
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)
    p = _safe_ws(str(body.get("ws") or ""))
    if not p:
        return JSONResponse({"ok": False, "message": "工作目录不存在"}, status_code=400)
    action = body.get("action")
    path = str(body.get("path") or "")

    if action == "init":
        code, out = await asyncio.to_thread(_git, ["init"], str(p))
        return {"ok": code == 0, "message": out.strip()[-200:] or "已初始化"}

    if action == "restore":
        if path:
            if ".." in path or path.startswith("/") or ":" in path:
                return JSONResponse({"ok": False, "message": "非法路径"}, status_code=400)
            # 先确认这个文件确实在改动列表里，避免误删
            _, porcelain = await asyncio.to_thread(_git, ["status", "--porcelain=v1"], str(p))
            if path not in porcelain:
                return JSONResponse({"ok": False, "message": "这个文件当前没有改动"}, status_code=400)
            code, out = await asyncio.to_thread(_git, ["checkout", "--", path], str(p))
            if code != 0:   # 新文件用 checkout 恢复不了，直接删
                code, out = await asyncio.to_thread(_git, ["clean", "-f", "--", path], str(p))
            return {"ok": code == 0, "message": (out or "已还原").strip()[-200:]}
        # 全部还原 —— ⚠️ 没有提交过的仓库绝不能执行：那时所有文件都是"未跟踪"，
        # clean -fd 会把它们全删掉（实测踩过）。必须先有存档点。
        code_h, _ = await asyncio.to_thread(_git, ["rev-parse", "--verify", "HEAD"], str(p))
        if code_h != 0:
            return JSONResponse({"ok": False, "message":
                "这个仓库还没有任何提交（没有存档点）。现在还原会把所有文件删掉，已阻止。"
                "先点「提交全部」存一个存档点，之后再有改动就能一键回滚了。"}, status_code=400)
        await asyncio.to_thread(_git, ["checkout", "--", "."], str(p))
        code, out = await asyncio.to_thread(_git, ["clean", "-fd"], str(p))
        print(f"[git] 全部还原：{p}")
        return {"ok": code == 0, "message": (out or "已全部还原").strip()[-200:]}

    if action == "commit":
        msg = str(body.get("message") or "").strip() or f"webui: 改动 {time.strftime('%m-%d %H:%M')}"
        await asyncio.to_thread(_git, ["add", "-A"], str(p))
        # 没配 git 身份时提交会直接失败（"please tell me who you are"），
        # 所以没配就临时用本机用户名兜一个，不去动他的全局配置
        code_n, _ = await asyncio.to_thread(_git, ["config", "user.name"], str(p))
        extra: List[str] = []
        if code_n != 0:
            uname = os.environ.get("USERNAME") or os.environ.get("USER") or Path.home().name or "local"
            extra = ["-c", f"user.name={uname}", "-c", f"user.email={uname}@localhost"]
            print(f"[git] 仓库没配身份，临时用 {uname}@localhost 提交")
        code, out = await asyncio.to_thread(_git, extra + ["commit", "-m", msg], str(p))
        return {"ok": code == 0, "message": out.strip()[-300:]}

    return JSONResponse({"ok": False, "message": f"不认识的动作：{action}"}, status_code=400)


@app.post("/api/attach-text")
async def api_attach_text(request: Request):
    """把一段长文本当成附件（前端粘贴超长文本时用，免得撑爆输入框）"""
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)
    text = body.get("text") or ""
    if not text.strip():
        return JSONResponse({"ok": False, "message": "内容是空的"}, status_code=400)
    name = str(body.get("name") or "").strip() or f"粘贴的文本_{time.strftime('%H%M%S')}.txt"
    if not name.lower().endswith((".txt", ".md", ".log", ".csv", ".json")):
        name += ".txt"
    meta = await _ingest(name, text.encode("utf-8"))
    return {"ok": True, "id": meta["id"], "name": meta["name"], "kind": meta["kind"],
            "size": meta["size"], "note": meta["note"], "text_chars": len(meta["text"])}


# ==========================================================================
#  记忆 & 终端会话：让网页端和终端共用同一个"大脑"
#  记忆本来就是共享的（各项目目录的 memory 都是同一份的目录链接），
#  这里只是把它显示出来、并让会话可以互相接管。
# ==========================================================================
CLAUDE_PROJECTS = Path.home() / ".claude" / "projects"
# 长期记忆目录：config.json 里配了 memory_dir 就用它（打包版指到自己的数据目录里），
# 否则用本机终端那一套（终端和网页共用同一份记忆）
MEM_DIR = (Path(CONFIG["memory_dir"]) if CONFIG["memory_dir"]
           else CLAUDE_PROJECTS / "C--WINDOWS-system32" / "memory")
with contextlib.suppress(Exception):
    MEM_DIR.mkdir(parents=True, exist_ok=True)


# --------------------------------------------------------------------------
#  工作目录 ≠ 失忆：claude 是按「工作目录」找记忆的
#  （~/.claude/projects/<目录名转义>/memory/），换个没用过的目录就是一份空白记忆。
#  这里把每个项目目录的记忆都接到共享记忆上——用户在任何目录里干活都认识他。
# --------------------------------------------------------------------------
def _proj_slug(cwd) -> str:
    """claude 的项目目录名：非字母数字一律换成 -（实测 E:\\git练习 → E--git--）"""
    return re.sub(r"[^A-Za-z0-9]", "-", str(cwd))


def _link_memory_to_shared(proj: Path) -> str:
    """把 proj/memory 接到共享记忆；幂等，返回给日志的一句话（没事发生就返回空串）"""
    mem = proj / "memory"
    try:
        if mem.is_dir() and mem.resolve() == MEM_DIR.resolve():
            return ""                       # 已经是共享那一份
        if mem.is_dir() and any(mem.iterdir()):
            # 里面已经有独立记忆 → **一律不动它，只记一笔日志**。
            # ⚠️ 这里原来做的是"把内容并进共享目录、原目录改名留档、再换成链接"，
            #    2026-09-13 实测踩了大坑：当 MEM_DIR 指向**别的位置**（打包版/测试实例配了
            #    自己的 memory_dir）时，它会把用户真正的记忆整个搬走——主人的记忆目录被挪到
            #    %LOCALAPPDATA%\Jingyu\memory 下面、原位只剩链接，清理那个测试目录就等于删记忆。
            #    原则改死：**只给空的补链接，绝不搬有东西的**。
            return f"跳过 {proj.name}：里面已有记忆，不搬动它"
        elif mem.exists():
            mem.rmdir()                      # 空目录：直接让位
        proj.mkdir(parents=True, exist_ok=True)
        r = subprocess.run(["cmd", "/c", "mklink", "/J", str(mem), str(MEM_DIR)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            return f"接记忆失败({proj.name})：{(r.stderr or r.stdout).strip()}"
        return f"已把 {proj.name} 的记忆接上共享记忆"
    except Exception as e:
        return f"接记忆出错({proj.name})：{e}"


def _sweep_memory_links() -> None:
    """启动时扫一遍：已有的项目目录全部接上共享记忆（「新目录失忆」的坑就此补上）"""
    if not MEM_DIR.is_dir() or not CLAUDE_PROJECTS.is_dir():
        return
    for proj in CLAUDE_PROJECTS.iterdir():
        if proj.is_dir():
            msg = _link_memory_to_shared(proj)
            if msg:
                print(f"[memory] {msg}")


_sweep_memory_links()


def _memory_text(max_chars: int = 6000) -> str:
    """取记忆索引（MEMORY.md），用于给 CHAT 模式当上下文"""
    f = MEM_DIR / "MEMORY.md"
    with contextlib.suppress(Exception):
        return f.read_text(encoding="utf-8", errors="replace")[:max_chars]
    return ""


@app.get("/api/memory")
async def api_memory():
    files = []
    if MEM_DIR.is_dir():
        for f in sorted(MEM_DIR.glob("*.md")):
            with contextlib.suppress(Exception):
                files.append({"name": f.name, "size": f.stat().st_size,
                              "mtime": f.stat().st_mtime})
    return {"ok": True, "dir": str(MEM_DIR), "count": len(files), "files": files}


@app.get("/api/memory/{name}")
async def api_memory_get(name: str):
    if not re.match(r"^[\w一-鿿.\-]{1,64}\.md$", name):
        return JSONResponse({"ok": False, "message": "非法文件名"}, status_code=400)
    f = MEM_DIR / name
    if not f.exists():
        return JSONResponse({"ok": False, "message": "没有这个记忆文件"}, status_code=404)
    return {"ok": True, "name": name, "content": f.read_text(encoding="utf-8", errors="replace")}


@app.post("/api/memory/{name}")
async def api_memory_save(name: str, request: Request):
    if not re.match(r"^[\w一-鿿.\-]{1,64}\.md$", name):
        return JSONResponse({"ok": False, "message": "非法文件名"}, status_code=400)
    body = await request.json()
    content = body.get("content")
    if not isinstance(content, str):
        return JSONResponse({"ok": False, "message": "content 必须是字符串"}, status_code=400)
    if len(content) > 200_000:
        return JSONResponse({"ok": False, "message": "太大了"}, status_code=413)
    (MEM_DIR / name).write_text(content, encoding="utf-8")
    print(f"[memory] 已保存 {name}（{len(content)} 字）")
    return {"ok": True, "name": name, "bytes": len(content)}


def _session_meta(path: Path) -> Optional[Dict[str, Any]]:
    """只读每个会话文件的前几行，拿到 cwd / 首条消息 / 时间（快）"""
    try:
        cwd, first_user, ts, n_lines = "", "", 0.0, 0
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f):
                n_lines = i + 1
                if i > 60:
                    break
                with contextlib.suppress(Exception):
                    d = json.loads(line)
                    if not cwd and d.get("cwd"):
                        cwd = d["cwd"]
                    if not ts and d.get("timestamp"):
                        ts = time.mktime(time.strptime(d["timestamp"][:19], "%Y-%m-%dT%H:%M:%S"))
                    if not first_user and d.get("type") == "user":
                        m = d.get("message") or {}
                        c = m.get("content")
                        if isinstance(c, str):
                            first_user = c
                        elif isinstance(c, list):
                            first_user = " ".join(x.get("text", "") for x in c
                                                  if isinstance(x, dict) and x.get("type") == "text")
        st = path.stat()
        return {"id": path.stem, "project": path.parent.name, "cwd": cwd,
                "title": (first_user or "(空会话)").replace("\n", " ")[:70],
                "ts": ts or st.st_mtime, "size": st.st_size, "lines": n_lines}
    except Exception:
        return None


@app.get("/api/claude-sessions")
async def api_claude_sessions(limit: int = 80):
    """列出终端里跑过的 Claude Code 会话（读 ~/.claude/projects/*/*.jsonl）"""
    items: List[Dict[str, Any]] = []
    if CLAUDE_PROJECTS.is_dir():
        for proj in CLAUDE_PROJECTS.iterdir():
            if not proj.is_dir():
                continue
            for f in proj.glob("*.jsonl"):
                m = await asyncio.to_thread(_session_meta, f)
                if m:
                    items.append(m)
    items.sort(key=lambda x: -x["ts"])
    return {"ok": True, "items": items[:limit], "total": len(items)}


@app.get("/api/claude-sessions/{sid}")
async def api_claude_session_read(sid: str, max_msgs: int = 200):
    """把某个终端会话的完整对话读出来（转成我们前端的格式，好显示在网页上）"""
    if not re.match(r"^[A-Za-z0-9_-]{8,64}$", sid):
        return JSONResponse({"ok": False, "message": "非法 session id"}, status_code=400)
    hit = None
    for proj in CLAUDE_PROJECTS.iterdir() if CLAUDE_PROJECTS.is_dir() else []:
        f = proj / f"{sid}.jsonl"
        if f.exists():
            hit = f
            break
    if not hit:
        return JSONResponse({"ok": False, "message": "找不到这个会话"}, status_code=404)

    msgs, cwd = [], ""
    with contextlib.suppress(Exception):
        with open(hit, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                with contextlib.suppress(Exception):
                    d = json.loads(line)
                    if not cwd and d.get("cwd"):
                        cwd = d["cwd"]
                    t = d.get("type")
                    m = d.get("message") or {}
                    if t == "user":
                        c = m.get("content")
                        if isinstance(c, str):
                            txt = c
                        elif isinstance(c, list):
                            txt = " ".join(x.get("text", "") for x in c
                                           if isinstance(x, dict) and x.get("type") == "text")
                        else:
                            txt = ""
                        txt = (txt or "").strip()
                        # 跳过工具回执和系统注入的噪声
                        if txt and not txt.startswith("<") and "system-reminder" not in txt[:40]:
                            msgs.append({"id": uid_hex(), "role": "user", "ts": 0, "content": txt})
                    elif t == "assistant":
                        c = m.get("content") or []
                        txt = " ".join(x.get("text", "") for x in c
                                       if isinstance(x, dict) and x.get("type") == "text").strip()
                        if txt:
                            msgs.append({"id": uid_hex(), "role": "assistant", "ts": 0, "content": txt})
    msgs = msgs[-max_msgs:]
    return {"ok": True, "id": sid, "cwd": cwd, "project": hit.parent.name,
            "messages": msgs, "count": len(msgs)}


def uid_hex() -> str:
    return os.urandom(6).hex()


async def _ingest(name: str, raw: bytes) -> Dict[str, Any]:
    """把一份文件落盘并解析，返回给前端的元信息"""
    display = os.path.basename(name) or "file"
    ext = display.rsplit(".", 1)[-1].lower() if "." in display else ""
    kind = _kind_of(ext)
    fid = f"{int(time.time()*1000):x}-{os.urandom(3).hex()}"
    safe = re.sub(r"[^\w一-鿿.\-]", "_", display)[:70]
    stored = FILES_DIR / f"{fid}_{safe}"
    stored.write_bytes(raw)

    meta: Dict[str, Any] = {"id": fid, "name": display, "ext": ext, "kind": kind,
                            "size": len(raw), "stored": stored.name,
                            "images": [], "text": "", "note": ""}
    if kind == "image":
        img_name = f"{fid}.{ext if ext in ('jpg', 'jpeg', 'png', 'webp', 'gif') else 'png'}"
        (IMG_DIR / img_name).write_bytes(raw)
        meta["images"] = [img_name]
        meta["note"] = "图片直接交给模型看"
    else:
        try:
            res = await asyncio.to_thread(_parse_attachment, stored, kind)
            meta["text"] = res.get("text") or ""
            meta["images"] = res.get("images") or []
            meta["note"] = res.get("note") or ""
        except Exception as e:
            meta["note"] = f"解析出错：{type(e).__name__}: {e}"
    (FILES_DIR / f"{fid}.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    print(f"[attach] {display} ({kind}) {len(raw)}B -> {meta['note']}")
    return meta


@app.post("/api/attach")
async def api_attach(request: Request, name: str = "file"):
    raw = await request.body()
    if not raw:
        return JSONResponse({"ok": False, "message": "空文件"}, status_code=400)
    if len(raw) > 512 * 1024 * 1024:
        return JSONResponse({"ok": False, "message": "文件太大了（>512MB）"}, status_code=413)
    meta = await _ingest(name, raw)
    return {"ok": True, "id": meta["id"], "name": meta["name"], "kind": meta["kind"],
            "size": meta["size"], "note": meta["note"],
            "text_chars": len(meta["text"]), "images": meta["images"]}


@app.post("/api/attach-path")
async def api_attach_path(request: Request):
    """给 /attach 斜杠命令用：直接按本机路径加附件（不经过浏览器上传）"""
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "请求体不是合法 JSON"}, status_code=400)
    p = Path(str(body.get("path") or "").strip().strip('"'))
    if not p.is_file():
        return JSONResponse({"ok": False, "message": f"找不到文件：{p}"}, status_code=404)
    try:
        raw = p.read_bytes()
    except Exception as e:
        return JSONResponse({"ok": False, "message": f"读不了：{e}"}, status_code=500)
    meta = await _ingest(p.name, raw)
    return {"ok": True, "id": meta["id"], "name": meta["name"], "kind": meta["kind"],
            "size": meta["size"], "note": meta["note"]}


def _load_attachment(fid: str) -> Optional[Dict[str, Any]]:
    if not re.match(r"^[A-Za-z0-9_-]{1,64}$", fid or ""):
        return None
    f = FILES_DIR / f"{fid}.json"
    if not f.exists():
        return None
    with contextlib.suppress(Exception):
        return json.loads(f.read_text(encoding="utf-8"))
    return None


# ------------------------ 对话备份(存本地文件夹) ------------------------
@app.post("/api/conversations/save")
async def conv_save(request: Request):
    try:
        body = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "message": "不是合法 JSON"}, status_code=400)
    cid = str(body.get("id") or "")
    data = body.get("data")
    if not SAFE_ID.match(cid) or not isinstance(data, dict):
        return JSONResponse({"ok": False, "message": "id 或 data 不合法"}, status_code=400)
    tmp = CONV_DIR / (cid + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(CONV_DIR / (cid + ".json"))     # 原子替换,断电不会写坏
    return {"ok": True}


@app.get("/api/conversations/list")
async def conv_list():
    items = []
    for f in CONV_DIR.glob("*.json"):
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
            items.append({
                "id": d.get("id") or f.stem,
                "title": d.get("title") or "(无标题)",
                "updatedAt": d.get("updatedAt") or 0,
                "count": len(d.get("messages") or []),
            })
        except Exception:
            continue
    items.sort(key=lambda x: x.get("updatedAt", 0), reverse=True)
    return {"ok": True, "items": items}


@app.get("/api/conversations/{cid}")
async def conv_get(cid: str):
    if not SAFE_ID.match(cid):
        return JSONResponse({"ok": False, "message": "非法 id"}, status_code=400)
    f = CONV_DIR / (cid + ".json")
    if not f.exists():
        return JSONResponse({"ok": False, "message": "没有这个对话"}, status_code=404)
    try:
        return {"ok": True, "data": json.loads(f.read_text(encoding="utf-8"))}
    except Exception as e:
        return JSONResponse({"ok": False, "message": f"文件坏了: {e}"}, status_code=500)


@app.delete("/api/conversations/{cid}")
async def conv_delete(cid: str):
    if not SAFE_ID.match(cid):
        return JSONResponse({"ok": False, "message": "非法 id"}, status_code=400)
    f = CONV_DIR / (cid + ".json")
    with contextlib.suppress(Exception):
        f.unlink(missing_ok=True)
    return {"ok": True}


@app.post("/api/open-folder")
async def open_folder(which: str = "data"):
    """在资源管理器里打开数据文件夹,方便你直接看备份。"""
    target = DATA_DIR if which != "images" else IMG_DIR
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(target))          # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            os.system(f'open "{target}"')
        else:
            os.system(f'xdg-open "{target}"')
        return {"ok": True, "path": str(target)}
    except Exception as e:
        return {"ok": False, "message": f"打不开: {e}", "path": str(target)}


# --------------------------------------------------------------------------
# 启动
# --------------------------------------------------------------------------
def pick_port(host: str, port: int, tries: int = 12) -> int:
    """端口被占用就自动往后找,最多试 12 个。"""
    for i in range(tries):
        p = port + i
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            s.bind((host, p))
            s.close()
            return p
        except OSError:
            s.close()
            print(f"[port] {p} 被占用了,试下一个…")
    raise SystemExit(f"[x] {port}~{port + tries} 都被占用了,用 --port 换一个吧。")


def main() -> None:
    ap = argparse.ArgumentParser(description="DeepSeek 本地聊天 UI")
    ap.add_argument("--host", default=CONFIG.get("host", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(CONFIG.get("port", 80)))
    ap.add_argument("--no-browser", action="store_true", help="启动后不自动开浏览器")
    ap.add_argument("--reload", action="store_true", help="改代码自动重启(开发用)")
    args = ap.parse_args()

    port = pick_port(args.host, args.port)
    url = f"http://{args.host}:{port}/"

    # 把实际用的端口记下来：start.bat 靠它判断"服务是不是已经在跑了"
    with contextlib.suppress(Exception):
        (DATA_DIR / "port.txt").write_text(str(port), encoding="utf-8")

    if not args.no_browser:
        def _open():
            with contextlib.suppress(Exception):
                webbrowser.open(url)
        threading.Timer(1.2, _open).start()

    print(f"\n  在浏览器打开: {url}")
    print(f"  数据文件夹  : {DATA_DIR}\n")
    uvicorn.run("server:app" if args.reload else app,
                host=args.host, port=port,
                log_level="warning", reload=args.reload)


if __name__ == "__main__":
    main()
